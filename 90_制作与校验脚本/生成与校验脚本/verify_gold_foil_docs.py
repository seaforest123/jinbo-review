from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "docs" / "gold-foil-heritage"
PRD_PATH = DOC_DIR / "金箔非遗小程序产品需求文档PRD-V0.3.docx"
TASK_PATH = DOC_DIR / "金箔非遗项目四岗位7周任务与验收表V0.1.docx"


def all_text(doc: Document) -> str:
    blocks = [paragraph.text for paragraph in doc.paragraphs]
    blocks.extend(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(blocks)


def assert_zip_ok(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
    assert bad is None, f"Corrupt ZIP member in {path.name}: {bad}"


def verify_prd() -> dict[str, int | str]:
    assert PRD_PATH.exists() and PRD_PATH.stat().st_size > 50_000
    assert_zip_ok(PRD_PATH)
    doc = Document(PRD_PATH)
    text = all_text(doc)

    assert doc.core_properties.title == "金箔非遗小程序产品需求文档 PRD V0.3"
    assert "产品需求文档 PRD V0.3" in doc.sections[0].header.paragraphs[0].text
    for required in (
        "PRD V0.3",
        "状态／下次复核",
        "内部临时技术占位路线",
        "正式拍摄",
        "体验版发布",
        "四岗位7周任务与验收表V0.1",
        "研究补偿遵循知情自愿并与传播激励分开",
    ):
        assert required in text, f"Missing PRD requirement: {required}"
    assert "将本PRD版本更新为V0.2" not in text

    functional_ids = {
        match.group(0)
        for match in re.finditer(
            r"(?:HOME|EXP|INT|RES|ARC|ACT|AR|SHR)-\d{2}", text
        )
    }
    assert len(functional_ids) == 28, functional_ids

    risk_headers = [
        "风险",
        "影响域／等级",
        "触发信号",
        "应对预案",
        "状态／下次复核",
        "责任人",
    ]
    assert any(
        [cell.text.strip() for cell in table.rows[0].cells] == risk_headers
        for table in doc.tables
    )
    return {
        "functional_ids": len(functional_ids),
        "tables": len(doc.tables),
        "bytes": PRD_PATH.stat().st_size,
    }


def verify_task_table() -> dict[str, int | str]:
    assert TASK_PATH.exists() and TASK_PATH.stat().st_size > 40_000
    assert_zip_ok(TASK_PATH)
    doc = Document(TASK_PATH)
    text = all_text(doc)

    assert doc.core_properties.title == "金箔非遗项目四岗位7周任务与验收表 V0.1"
    assert doc.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert "四岗位7周任务与验收表 V0.1" in doc.sections[0].header.paragraphs[0].text
    assert "P0最终验收与交接清单" in text.replace(" ", "")

    weekly_headers = [
        "岗位",
        "主责任务",
        "协作／前置依赖",
        "本周交付物",
        "验收标准",
        "验收人",
        "状态",
    ]
    weekly_tables = [
        table
        for table in doc.tables
        if [cell.text.strip() for cell in table.rows[0].cells] == weekly_headers
    ]
    assert len(weekly_tables) == 7
    assert sum(len(table.rows) - 1 for table in weekly_tables) == 28
    assert all(
        row.cells[-1].text.strip() == "未开始"
        for table in weekly_tables
        for row in table.rows[1:]
    )
    for week in range(1, 8):
        assert f"第{week}周｜" in text
    for role in ("1 产品与前端", "2 3D与技术", "3 内容与档案", "4 调研与实践"):
        assert sum(role in row.cells[0].text for table in weekly_tables for row in table.rows[1:]) == 7

    forbidden = ("学生A", "学生B", "学生C", "学生D", "待填写姓名")
    assert not any(item in text for item in forbidden)
    return {
        "weeks": len(weekly_tables),
        "role_week_rows": sum(len(table.rows) - 1 for table in weekly_tables),
        "tables": len(doc.tables),
        "bytes": TASK_PATH.stat().st_size,
    }


def main() -> None:
    prd = verify_prd()
    task = verify_task_table()
    print(f"PRD_OK {prd}")
    print(f"TASK_OK {task}")
    print("ALL_DOCUMENT_CHECKS_PASSED")


if __name__ == "__main__":
    main()
