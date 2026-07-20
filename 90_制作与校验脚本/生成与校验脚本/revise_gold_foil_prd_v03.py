from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_gold_foil_project_brief_docx import (  # noqa: E402
    GOLD_DARK,
    INK,
    MUTED,
    add_table,
    apply_numbering,
    create_numbering,
    fill_cell,
    set_run_font,
)


SOURCE = Path(r"C:\Users\Andrew\Downloads\金箔非遗小程序产品需求文档PRD-V0_2.docx")
OUT = ROOT / "docs" / "gold-foil-heritage" / "金箔非遗小程序产品需求文档PRD-V0.3.docx"


def replace_in_runs(paragraph, old: str, new: str, required: bool = True) -> bool:
    replaced = False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            replaced = True
    if replaced:
        return True

    full = "".join(run.text for run in paragraph.runs)
    if old in full and paragraph.runs:
        paragraph.runs[0].text = full.replace(old, new)
        for run in paragraph.runs[1:]:
            run.text = ""
        return True

    if required:
        raise AssertionError(f"Expected text not found in paragraph: {old}")
    return False


def find_paragraph(doc: Document, startswith: str):
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(startswith)]
    if len(matches) != 1:
        raise AssertionError(f"Expected one paragraph starting with {startswith!r}, found {len(matches)}")
    return matches[0]


def set_labeled_paragraph(paragraph, label: str, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.paragraph_format.line_spacing = 1.25
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=10.5, bold=True, color=GOLD_DARK)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.5, color=INK)


def find_table(doc: Document, headers: list[str]):
    for table in doc.tables:
        values = [cell.text.strip() for cell in table.rows[0].cells]
        if values == headers:
            return table
    raise AssertionError(f"Table not found: {headers}")


def remove_num_pr(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def set_table_indent(table, value: int) -> None:
    """Align the visible table border with its existing cell start margin."""
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(value))
    tbl_ind.set(qn("w:type"), "dxa")


def mark_row_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def normalize_existing_table_metadata(doc: Document) -> None:
    # The one-row project metric strip functions as a semantic label row.
    mark_row_as_header(doc.tables[1].rows[0])
    # The two appended V0.2 tables use 100-DXA cell starts; match their table indent.
    set_table_indent(doc.tables[-2], 100)
    set_table_indent(doc.tables[-1], 100)


def replace_risk_table(doc: Document) -> None:
    old = find_table(doc, ["风险", "影响域／等级", "触发信号", "应对预案", "责任人"])
    rows = [
        (
            "首次专业访谈延期，未能48小时内锁定制箔／贴箔",
            "内容／进度｜高",
            "访谈未如期安排；专业意见尚不足以锁定主线",
            "设置访谈硬截止日（不晚于7月23日）；逾期时仅建立内部临时技术占位路线，用于交互和性能验证，不对外发布；专业确认后可以调整或推翻临时路线。",
            "观察中／7月20日",
            "指导教师",
        ),
        (
            "传承人拍摄、形象或工艺授权未能及时取得",
            "内容／合规｜高",
            "授权协商超过2周仍未达成",
            "提前准备书面授权模板；未授权素材仅用于内部联调，不进入公开版本；必要时使用已获授权的专业老师讲解素材作为内部P0替代。",
            "观察中／7月24日",
            "岗位3",
        ),
        (
            "XR-Frame手势或动作识别技术不可行",
            "技术｜高",
            "单步骤原型仍无法稳定识别",
            "降级为点击、拖动或滑动等简化交互，保留引导完成原则；在P0开发前暴露并记录降级决定。",
            "验证中／8月6日",
            "岗位2",
        ),
        (
            "中低端机型加载时间或资源体积超预算",
            "技术／性能｜中高",
            "中低档机加载超过15秒或帧率过低",
            "先压缩模型与纹理，再降低特效和交互复杂度；仍不达标时启用2D、图文或视频替代路径。",
            "验证中／8月6日",
            "岗位2",
        ),
        (
            "800+真实用户目标未达成",
            "增长／指标｜中",
            "30—50人测试后，渠道转化预估明显低于目标",
            "按渠道执行获客路径；申报前根据真实测试校准目标，数据不足时如实调整，不将目标值表述为既有成果。",
            "未触发／8月20日",
            "岗位4",
        ),
        (
            "团队成员退出或投入不足",
            "团队｜高",
            "原型评审或P0评审交付延迟",
            "明确岗位备份人；代码、素材、数据和文档每周交接，不在临近交付时集中补齐。",
            "未触发／每周五",
            "指导教师",
        ),
        (
            "研究样本知情同意或招募不足",
            "研究｜中",
            "首轮正式研究招募未达到计划样本",
            "结合线下活动招募；研究补偿遵循知情自愿并与传播激励分开，不以奖励替代告知和明确同意。",
            "未触发／8月13日",
            "岗位4",
        ),
        (
            "内容审核形成瓶颈",
            "内容／进度｜中",
            "待审内容积压超过约定审核周期",
            "约定3个工作日反馈的审核目标，批量提交并保留版本；未审核内容不得进入公开版本。",
            "未触发／8月13日",
            "岗位3",
        ),
    ]

    new = add_table(
        doc,
        ["风险", "影响域／等级", "触发信号", "应对预案", "状态／下次复核", "责任人"],
        rows,
        [1.55, 1.0, 1.35, 2.15, 1.15, 0.8],
        font_size=7.2,
    )
    for row in new.rows[1:]:
        for idx in (0, 2, 3, 4):
            text = row.cells[idx].text
            fill_cell(row.cells[idx], text, size=7.2, align=WD_ALIGN_PARAGRAPH.LEFT)

    old._tbl.addprevious(new._tbl)
    old._tbl.getparent().remove(old._tbl)

    trailing = doc.paragraphs[-1]
    if not trailing.text.strip():
        trailing._p.getparent().remove(trailing._p)

    usage = find_paragraph(doc, "使用说明｜新增风险")
    set_labeled_paragraph(
        usage,
        "使用说明｜",
        "新增风险须补充影响域、触发信号、应对预案、当前状态和复核时间；每周更新状态及实际处置记录。已缓解风险保留记录并标注状态，不删除。",
    )


def restart_acquisition_numbering(doc: Document) -> None:
    paragraphs = [
        find_paragraph(doc, "在home_cta_click事件的source字段中"),
        find_paragraph(doc, "3场线下活动应作为获客重点抓手"),
        find_paragraph(doc, "短视频发布节奏建议与活动时间"),
    ]
    set_labeled_paragraph(
        paragraphs[1],
        "",
        "3场线下活动应作为获客重点抓手：现场直接扫码体验；如设置传播激励，应保持自愿、适度且与正式研究参与无关。研究样本的知情同意与研究补偿另行管理。",
    )
    decimal_id = create_numbering(doc, "decimal")
    for paragraph in paragraphs:
        remove_num_pr(paragraph)
        apply_numbering(paragraph, decimal_id)


def revise_schedule(doc: Document) -> None:
    table = find_table(doc, ["周次", "日期区间", "里程碑", "关键产出／验收对照"])
    updates = {
        "第2周": (
            "低保真流程原型、内容拍摄清单与授权边界确认；3D单步骤技术原型启动",
            "岗位1原型、岗位3拍摄与授权清单；小程序名称和视觉识别锁定",
        ),
        "第3周": (
            "3D单步骤原型完成中低档机测试；正式拍摄；前端骨架、云环境和埋点字典完成；P0开发启动",
            "模型样例与6台设备记录；开场及三步骤原始素材；埋点字典初稿",
        ),
        "第4周": (
            "P0核心流程开发；视频剪辑、字幕、专业审核与内容导入",
            "HOME／EXP／INT／RES逐项开发；开场和3段步骤视频审核通过；问卷与同意文本定稿",
        ),
        "第5周": (
            "P0联调、体验版发布、埋点核验、6台真机兼容测试和30—50人测试",
            "体验版二维码；真机与弱网／权限记录；测试问题清单和数据摘要",
        ),
        "第6周": (
            "问题修复、性能与降级优化、内容导入收尾和P0演示数据整理",
            "完成率漏斗、退出步骤统计、降级测试记录；内容台账审核状态清零",
        ),
        "第7周": (
            "版本冻结、最终演练、申报材料和演示准备",
            "P0验收通过；材料、PPT、演示脚本和部署备份定稿；备赛9月初比赛",
        ),
    }
    for row in table.rows[1:]:
        week = row.cells[0].text.strip()
        if week in updates:
            fill_cell(row.cells[2], updates[week][0], size=8.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            fill_cell(row.cells[3], updates[week][1], size=8.2, align=WD_ALIGN_PARAGRAPH.LEFT)

    principle = find_paragraph(doc, "排期原则｜P0申报演示验收")
    set_labeled_paragraph(
        principle,
        "排期原则｜",
        "P0申报演示验收须在参赛前至少预留1周缓冲；第5周测试发现重大问题时，第6周优先修复。三步流程必须完整，至少一个步骤达到高保真；技术风险触发时，其余步骤可采用简化交互，但不得取消真实视频、即时反馈和结果页。岗位周任务详见《金箔非遗项目四岗位7周任务与验收表V0.1》。",
    )


def compact_schedule_and_appendix(doc: Document) -> None:
    """Keep the short version note with the appendix instead of on a blank tail page."""
    schedule = find_table(doc, ["周次", "日期区间", "里程碑", "关键产出／验收对照"])
    for row in schedule.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=7.8, color=INK)

    for prefix in (
        "微信XR-Frame官方文档",
        "腾讯云开发CloudBase官方文档",
        "《中华人民共和国个人信息保护法》",
    ):
        paragraph = find_paragraph(doc, prefix)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            set_run_font(run, size=9.2, color=INK)

    version_note = find_paragraph(doc, "版本说明｜PRD V0.3")
    version_note.paragraph_format.line_spacing = 1.0
    version_note.paragraph_format.space_before = Pt(2)
    version_note.paragraph_format.space_after = Pt(0)
    for index, run in enumerate(version_note.runs):
        set_run_font(run, size=9.2, bold=index == 0, color=GOLD_DARK if index == 0 else INK)


def revise_versions(doc: Document) -> None:
    subtitle = find_paragraph(doc, "PRD V0.2｜")
    replace_in_runs(subtitle, "PRD V0.2", "PRD V0.3")

    header = doc.sections[0].header.paragraphs[0]
    replace_in_runs(header, "PRD V0.1", "PRD V0.3")

    stale = find_paragraph(doc, "工艺路线锁定后")
    replace_in_runs(
        stale,
        "并将本PRD版本更新为V0.2",
        "并将本PRD更新为下一正式版本（预计V0.4）",
    )

    version_note = find_paragraph(doc, "版本说明｜PRD V0.2")
    replace_in_runs(version_note, "PRD V0.2", "PRD V0.3")

    doc.core_properties.title = "金箔非遗小程序产品需求文档 PRD V0.3"
    doc.core_properties.subject = "广东省攀登计划金箔非遗数字体验项目产品需求与执行基线"


def validate(doc: Document) -> None:
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    required = [
        "PRD V0.3",
        "状态／下次复核",
        "内部临时技术占位路线",
        "正式拍摄",
        "体验版发布",
        "四岗位7周任务与验收表V0.1",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise AssertionError(f"Missing required content: {missing}")
    if "将本PRD版本更新为V0.2" in text:
        raise AssertionError("Stale V0.2 update sentence remains")
    functional_ids = {
        cell.text.strip()
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip().startswith(("HOME-", "EXP-", "INT-", "RES-", "ARC-", "ACT-", "AR-", "SHR-"))
    }
    if len(functional_ids) != 28:
        raise AssertionError(f"Expected 28 functional IDs, found {len(functional_ids)}")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    revise_versions(doc)
    replace_risk_table(doc)
    restart_acquisition_numbering(doc)
    revise_schedule(doc)
    compact_schedule_and_appendix(doc)
    normalize_existing_table_metadata(doc)
    validate(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
