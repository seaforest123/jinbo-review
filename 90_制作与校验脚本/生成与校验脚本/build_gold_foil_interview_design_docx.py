from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_gold_foil_project_brief_docx import (  # noqa: E402
    CREAM_LIGHT,
    GOLD,
    GOLD_DARK,
    INK,
    MUTED,
    add_field,
    set_run_font,
    set_style_font,
)
from build_gold_foil_week1_kickoff_docx import (  # noqa: E402
    add_callout,
    add_table,
)


SOURCE = (
    ROOT
    / "docs"
    / "superpowers"
    / "specs"
    / "2026-07-19-gold-foil-first-interview-materials-pack-design.md"
)
OUTPUT = (
    ROOT
    / "docs"
    / "gold-foil-heritage"
    / "金箔非遗项目首次专业访谈材料包设计说明V0.1.docx"
)

PALE_GOLD = "F6F0E3"
SOFT_GOLD = "E4D3A5"


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_border_bottom(paragraph, color: str, size: int = 14) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.5, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 16, INK, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13, GOLD_DARK, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 11.5, GOLD_DARK, bold=True)
    h3.paragraph_format.space_before = Pt(9)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    clear_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(
        header.add_run("金箔非遗数字体验与活化传播项目"),
        size=8.2,
        bold=True,
        color=MUTED,
    )
    set_run_font(header.add_run("    |    首次专业访谈材料包设计说明 V0.1"), size=8.2, color=MUTED)

    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("第 "), size=8.5, color=MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.add_run(" 页"), size=8.5, color=MUTED)

    doc.core_properties.title = "金箔非遗项目首次专业访谈材料包设计说明 V0.1"
    doc.core_properties.subject = "首次专业访谈材料包的文件结构、边界、流程与验收标准"
    doc.core_properties.author = "金箔非遗项目团队"
    doc.core_properties.keywords = "金箔, 非遗, 访谈, 材料包, 纸面原型, 授权, 攀登计划"


class NumberingManager:
    def __init__(self, doc: Document):
        self.numbering = doc.part.numbering_part.element

    def _next_id(self, tag: str, attr: str) -> int:
        values = []
        for element in self.numbering.findall(qn(tag)):
            value = element.get(qn(attr))
            if value is not None and value.isdigit():
                values.append(int(value))
        return max(values, default=0) + 1

    def create(self, kind: str) -> int:
        abstract_id = self._next_id("w:abstractNum", "w:abstractNumId")
        num_id = self._next_id("w:num", "w:numId")

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)

        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if kind == "number" else "bullet")
        level.append(num_fmt)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "%1." if kind == "number" else "•")
        level.append(level_text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        level.append(suffix)

        p_pr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        level.append(p_pr)

        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Arial")
        r_fonts.set(qn("w:hAnsi"), "Arial")
        r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r_pr.append(r_fonts)
        level.append(r_pr)
        abstract.append(level)
        first_num = self.numbering.find(qn("w:num"))
        if first_num is None:
            self.numbering.append(abstract)
        else:
            self.numbering.insert(self.numbering.index(first_num), abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        self.numbering.append(num)
        return num_id

    @staticmethod
    def apply(paragraph, num_id: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num)


INLINE_CODE = re.compile(r"(`[^`]+`)")


def set_named_run_font(run, name: str, *, size: float, color: str, bold=False) -> None:
    set_run_font(run, size=size, color=color, bold=bold)
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)


def add_inline_text(paragraph, text: str, *, size=10.5, bold=False, color=INK) -> None:
    for part in INLINE_CODE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_named_run_font(run, "Consolas", size=9.2, color=GOLD_DARK, bold=True)
        else:
            set_run_font(paragraph.add_run(part), size=size, bold=bold, color=color)


def add_body(doc: Document, text: str, *, bold=False, color=INK, after=6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    add_inline_text(paragraph, text, bold=bold, color=color)
    return paragraph


def add_list_item(doc: Document, text: str, num_id: int, *, compact=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2 if compact else 4)
    paragraph.paragraph_format.line_spacing = 1.15 if compact else 1.25
    NumberingManager.apply(paragraph, num_id)
    add_inline_text(paragraph, text)
    return paragraph


def add_label(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(4)
    add_inline_text(paragraph, text, bold=True, color=GOLD_DARK)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_code_block(doc: Document, lines: list[str]) -> None:
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(10)
        paragraph.paragraph_format.right_indent = Pt(8)
        paragraph.paragraph_format.space_before = Pt(0 if index else 4)
        paragraph.paragraph_format.space_after = Pt(0 if index < len(lines) - 1 else 8)
        paragraph.paragraph_format.line_spacing = 1.05
        from build_gold_foil_week1_kickoff_docx import set_paragraph_shading

        set_paragraph_shading(paragraph, CREAM_LIGHT)
        set_named_run_font(
            paragraph.add_run(line or " "),
            "Microsoft YaHei",
            size=8.5,
            color=INK,
        )


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(5)
    set_run_font(kicker.add_run("广东省攀登计划 · 项目执行设计文件"), size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("金箔非遗项目首次专业访谈材料包"), size=24, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    set_run_font(subtitle.add_run("设计说明 V0.1｜文件结构、使用边界、执行流程与验收标准"), size=11.5, color=MUTED)
    set_paragraph_border_bottom(subtitle, GOLD, 16)

    add_table(
        doc,
        ["日期", "状态", "适用阶段", "交付结构"],
        [("2026年7月19日", "已确认", "首次专业访谈全流程", "9份独立文件＋1份总目录")],
        [1.3, 1.1, 2.2, 2.4],
        font_size=8.9,
        center_cols=(0, 1, 2, 3),
        zebra=False,
        min_row_mm=10,
    )
    add_callout(
        doc,
        "设计结论｜",
        "在尚无可运行Demo的情况下，以五屏纸面体验原型作为访谈核心工具；制箔与贴箔均保持为候选路线，所有工艺事实和公开范围必须经专业老师确认。",
        fill=PALE_GOLD,
    )
    add_body(doc, "本文件用于指导后续10份访谈材料的制作、使用和验收。", color=MUTED, after=9)


LABELS = {
    "材料包同时承担以下功能：",
    "主要内容：",
    "纸面原型包含五个界面：",
    "结构：",
    "授权项目分别设置“同意、不同意、需发布前复核”：",
    "包含四类模板：",
    "主要字段：",
    "检查范围：",
    "团队可以预填：",
    "必须留给访谈对象或现场确认：",
    "材料存放于独立文件夹，文件名统一为：",
    "示例：",
    "材料包建议目录：",
}


def parse_markdown(doc: Document, source: str) -> None:
    lines = source.splitlines()
    numbering = NumberingManager(doc)
    current_kind: str | None = None
    current_num_id: int | None = None
    in_code = False
    code_lines: list[str] = []
    current_section = ""

    body_started = False
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if not body_started:
            if stripped.startswith("## "):
                body_started = True
            else:
                continue

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
                current_kind = None
                current_num_id = None
            continue
        if in_code:
            code_lines.append(line)
            continue

        if not stripped:
            current_kind = None
            current_num_id = None
            continue

        if stripped.startswith("## "):
            current_section = stripped[3:]
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline_text(paragraph, stripped[3:], size=16, bold=True, color=INK)
            current_kind = None
            current_num_id = None
            continue
        if stripped.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline_text(paragraph, stripped[4:], size=13, bold=True, color=GOLD_DARK)
            current_kind = None
            current_num_id = None
            continue
        if stripped.startswith("#### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            add_inline_text(paragraph, stripped[5:], size=11.5, bold=True, color=GOLD_DARK)
            current_kind = None
            current_num_id = None
            continue

        if stripped.startswith("- "):
            if current_kind != "bullet":
                current_kind = "bullet"
                current_num_id = numbering.create("bullet")
            add_list_item(
                doc,
                stripped[2:],
                current_num_id,
                compact=current_section == "九、验收标准",
            )
            continue

        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if number_match:
            if current_kind != "number":
                current_kind = "number"
                current_num_id = numbering.create("number")
            add_list_item(
                doc,
                number_match.group(1),
                current_num_id,
                compact=current_section == "九、验收标准",
            )
            continue

        current_kind = None
        current_num_id = None
        if stripped in LABELS or (stripped.endswith("：") and len(stripped) <= 24):
            add_label(doc, stripped)
        elif stripped.startswith("`序号_材料名称"):
            add_callout(doc, "命名规则｜", stripped.strip("`"), fill=CREAM_LIGHT)
        else:
            add_body(doc, stripped)


def build_document(source_path: Path = SOURCE, output_path: Path = OUTPUT) -> Path:
    source = source_path.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    parse_markdown(doc, source)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    Document(output_path)
    return output_path


if __name__ == "__main__":
    print(build_document())
