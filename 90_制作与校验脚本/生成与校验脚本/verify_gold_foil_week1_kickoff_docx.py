from pathlib import Path

from docx import Document
from pypdf import PdfReader


DOCX = Path(
    r"D:\WeChatProjects\miniprogram-1\docs\gold-foil-heritage\金箔非遗项目第1周启动包V0.1.docx"
)
PDF = Path(
    r"D:\WeChatProjects\miniprogram-1\.tmp_week1_kickoff_render3\week1-kickoff-v01.pdf"
)


document = Document(DOCX)
text = "\n".join(
    [paragraph.text for paragraph in document.paragraphs]
    + [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
)

required_sections = [
    "项目启动会议程",
    "岗位志愿与团队分工",
    "首次专业访谈准备",
    "首次专业访谈提纲（上）",
    "首次专业访谈提纲（下）",
    "访谈结论与会后确认",
    "拍摄与授权确认清单",
    "第1周四岗位交付清单",
    "第1周验收记录",
    "决策与第2周行动日志",
    "进入第2周前确认",
]
roles = ["产品与前端", "3D与技术", "内容与档案", "调研与实践"]
missing = [item for item in required_sections + roles if item not in text]

question_numbers = {
    cell.text.strip()
    for table in document.tables
    for row in table.rows
    for cell in row.cells
    if cell.text.strip().isdigit()
}

assert not missing, f"missing={missing!r}"
assert {str(index) for index in range(1, 27)}.issubset(question_numbers)
assert not any(marker in text for marker in ("TODO", "TBD", "待补充占位"))
assert len(document.tables) == 28, f"tables={len(document.tables)}"

page_count = len(PdfReader(PDF).pages)
assert page_count == 12, f"pages={page_count}"
assert DOCX.stat().st_size > 50_000, f"size={DOCX.stat().st_size}"

print("content_check=PASS")
print("required_sections=11/11")
print("roles=4/4")
print("interview_questions=26/26")
print("tables=28")
print(f"rendered_pages={page_count}")
print(f"docx_bytes={DOCX.stat().st_size}")
