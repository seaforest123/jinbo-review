from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = (
    ROOT
    / "docs"
    / "gold-foil-heritage"
    / "金箔非遗项目首次专业访谈材料包设计说明V0.1.docx"
)


with ZipFile(DOCX) as package:
    assert package.testzip() is None
    numbering_root = ET.fromstring(package.read("word/numbering.xml"))

numbering_children = [child.tag.rsplit("}", 1)[-1] for child in numbering_root]
first_num_index = numbering_children.index("num")
assert "abstractNum" not in numbering_children[first_num_index:], (
    "Word numbering schema requires every abstractNum before the first num; "
    f"actual order tail={numbering_children[first_num_index:first_num_index + 6]!r}"
)

document = Document(DOCX)
paragraph_text = [paragraph.text.strip() for paragraph in document.paragraphs]
text = "\n".join(paragraph_text)

sections = [
    "一、设计目标",
    "二、已确认的总体原则",
    "三、文件组成",
    "四、预填与留空边界",
    "五、使用流程",
    "六、异常与风险处理",
    "七、文件组织与命名",
    "八、版式与可用性",
    "九、验收标准",
]
file_headings = [
    "00｜首次专业访谈材料包总目录与使用顺序",
    "01｜金箔非遗数字体验项目一页纸",
    "02｜三分钟轻量工艺体验纸面原型",
    "03｜纸面原型演示脚本与专家反馈表",
    "04｜首次专业访谈提纲与60分钟主持脚本",
    "05｜拍摄、录音及素材使用授权确认单",
    "06｜访谈对象身份与联系人信息确认表",
    "07｜时间地点确认及微信约访文本",
    "08｜现场设备与记录工具检查清单",
    "09｜访谈纪要、关键结论与会后确认单",
]
guards = [
    "不得把“制箔”或“贴箔”写成已经确定的首版路线",
    "授权采用逐项选择",
    "历史照片、作品图片和第三方材料需单独确认权利归属",
    "数字体验不得暗示用户已经掌握真实工艺",
    "未审核内容不得进入公开小程序",
    "序号_材料名称_V版本_日期.docx",
]
acceptance = [
    "10份文件均可独立打开、打印和填写",
    "总目录能够说明每份文件由谁使用、何时使用以及如何归档",
    "纸面原型不擅自确定制箔或贴箔路线",
    "所有待专业确认内容都有统一、醒目的标识",
    "授权范围可逐项选择，并区分内部研究与公开传播",
    "历史资料和第三方材料不被默认授权",
    "访谈提纲能够在60分钟内完成",
    "访谈纪要能够把结论转换为四岗位后续行动",
    "所有表格无文字截断、边界拥挤、跨页残行或不可填写区域",
    "全部文件完成结构审计、无障碍检查和逐页渲染检查后方可交付",
]

missing = [item for item in sections + file_headings + guards + acceptance if item not in text]
assert not missing, f"Missing required content: {missing!r}"
assert "待书面复核" not in text
assert "TODO" not in text and "TBD" not in text
assert len(document.tables) == 1
assert DOCX.stat().st_size > 45_000

heading_1 = sum(paragraph.style.name == "Heading 1" for paragraph in document.paragraphs)
heading_2 = sum(paragraph.style.name == "Heading 2" for paragraph in document.paragraphs)
assert heading_1 == 9, heading_1
assert heading_2 >= 13, heading_2

print("content_check=PASS")
print("design_sections=9/9")
print("material_definitions=10/10")
print("acceptance_criteria=10/10")
print(f"heading_1={heading_1}")
print(f"heading_2={heading_2}")
print(f"docx_bytes={DOCX.stat().st_size}")
