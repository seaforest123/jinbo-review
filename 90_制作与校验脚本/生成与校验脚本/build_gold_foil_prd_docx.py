from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_gold_foil_project_brief_docx import (  # noqa: E402
    CREAM,
    CREAM_LIGHT,
    GOLD,
    GOLD_DARK,
    INK,
    MUTED,
    SOFT_GRAY,
    WHITE,
    add_body,
    add_bullet,
    add_callout,
    add_heading,
    add_labeled_paragraph,
    add_metric_strip,
    add_numbered,
    add_table,
    configure_document,
    create_numbering,
    fill_cell,
    set_cell_shading,
    set_run_font,
)


OUT = ROOT / "docs" / "gold-foil-heritage" / "金箔非遗小程序产品需求文档PRD-V0.1.docx"


def replace_header(doc: Document) -> None:
    header_p = doc.sections[0].header.paragraphs[0]
    for run in list(header_p.runs):
        header_p._p.remove(run._r)
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    set_run_font(header_p.add_run("金箔非遗数字体验与活化传播项目"), size=8.5, bold=True, color=MUTED)
    set_run_font(header_p.add_run("    |    产品需求文档 PRD V0.1"), size=8.5, color=MUTED)


def emphasize_label_column(table) -> None:
    for ridx in range(1, len(table.rows)):
        cell = table.rows[ridx].cells[0]
        text = cell.text
        set_cell_shading(cell, CREAM)
        fill_cell(cell, text, bold=True, color=GOLD_DARK, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_requirement_table(doc, rows, weights=(1.0, 2.8, 2.2), font_size=8.8):
    return add_table(
        doc,
        ["编号", "产品需求", "验收口径"],
        rows,
        list(weights),
        font_size=font_size,
    )


def build_document() -> None:
    doc = Document()
    configure_document(doc)
    replace_header(doc)

    doc.core_properties.title = "金箔非遗小程序产品需求文档 PRD V0.1"
    doc.core_properties.subject = "广东省攀登计划金箔非遗数字体验项目产品需求"
    doc.core_properties.author = "项目团队"
    doc.core_properties.keywords = "金箔, 非遗, 微信小程序, 3D互动, XR-Frame, CloudBase, PRD"

    goals_bullets = create_numbering(doc, "bullet")
    nongoals_bullets = create_numbering(doc, "bullet")
    principles_bullets = create_numbering(doc, "bullet")
    flow_numbers = create_numbering(doc, "decimal")
    content_numbers = create_numbering(doc, "decimal")
    privacy_bullets = create_numbering(doc, "bullet")
    testing_bullets = create_numbering(doc, "bullet")
    decision_numbers = create_numbering(doc, "decimal")

    # memo_masthead, adapted to the approved gold project style.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(kicker.add_run("广东省攀登计划 · 产品执行文件"), size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("金箔非遗小程序产品需求文档"), size=26, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True
    set_run_font(subtitle.add_run("PRD V0.1｜3分钟轻量工艺体验与活化传播平台"), size=12.5, color=MUTED)

    metadata = add_table(
        doc,
        ["文档项", "当前定义"],
        [
            ("文档状态", "已确认需求基线；供原型、开发、内容、研究和测试共同执行"),
            ("核心用户", "广东高校学生，重点为18—24岁青年"),
            ("产品形态", "微信原生小程序＋CloudBase＋XR-Frame 3D互动"),
            ("首版主线", "3分钟互动体验；制箔或贴箔路线在首次专业访谈后锁定"),
            ("团队结构", "4名学生分任产品前端、3D技术、内容档案、调研实践"),
        ],
        [1.2, 4.8],
        zebra=False,
        font_size=9.4,
    )
    emphasize_label_column(metadata)

    add_callout(
        doc,
        "产品定位｜",
        "面向广东高校青年的金箔非遗轻量化数字体验与活化传播平台，以3D工艺互动为核心，连接传承人故事、数字档案与线下实践。",
    )
    add_callout(
        doc,
        "唯一待锁定决策｜",
        "首次访谈后，由传承人或专业老师确认首版主线采用“制箔”还是“贴箔”。锁定后只制作一条主线，不并行开发两套完整体验。",
        fill=SOFT_GRAY,
    )

    add_heading(doc, "文档使用说明", 1)
    add_body(
        doc,
        "本文件定义首版小程序的产品边界、用户流程、功能要求、数据与隐私规则、内容发布流程、技术约束和验收标准。产品原型、开发任务、内容采集和用户研究均应引用本文件中的需求编号。",
    )
    add_labeled_paragraph(doc, "优先级说明｜", "P0为申报演示必须完成；P1为正式首版应完成；P2为首版明确延后。")
    add_labeled_paragraph(doc, "变更规则｜", "新增需求必须说明其对应的项目目标、评价指标、负责人和对进度的影响，经指导教师与四岗位共同评审后方可纳入。")

    add_heading(doc, "1. 产品目标与成功标准", 1)
    add_heading(doc, "1.1 产品目标", 2)
    for item in [
        "让第一次接触金箔非遗的高校青年，在不注册、不先学习长篇知识的情况下愿意进入体验。",
        "用3分钟可完成的3D工艺互动建立对人物、工具、材料和工艺逻辑的初步认识。",
        "把传承人真实影像、经审核的数字档案和校园实践连接为可持续传播路径。",
        "通过行为数据、前后测、问卷、访谈和活动记录验证传播效果。",
    ]:
        add_bullet(doc, item, goals_bullets)

    add_heading(doc, "1.2 非目标", 2)
    for item in [
        "首版不建设独立APP、商城、支付、直播带货或复杂会员系统。",
        "首版不制作大型3D游戏、多关卡重度游戏化或完整空间AR。",
        "首版不同时覆盖多个非遗门类，也不同时制作制箔与贴箔两条完整主线。",
        "首版不以考试、排行榜、证书或强制分享作为核心留存手段。",
    ]:
        add_bullet(doc, item, nongoals_bullets)

    add_heading(doc, "1.3 成功指标", 2)
    add_metric_strip(doc, [("800+", "真实用户"), ("60%+", "体验完成率"), ("100+", "配对前后测"), ("3场", "校园实践")])
    add_table(
        doc,
        ["指标层", "指标定义", "目标值", "证据来源"],
        [
            ("触达", "去重真实用户数；校园体验或推广活动", "用户800+；活动3场", "匿名行为数据、活动记录"),
            ("体验", "开始核心体验后完成全部3步", "完成率不低于60%", "experience_start / complete"),
            ("深度", "完成后进入故事、档案或活动", "继续探索率不低于25%", "result_action、archive_view"),
            ("认知", "同一研究样本的前后测得分变化", "样本100+；均值提升20%目标", "研究编号关联的配对数据"),
            ("内容", "审核发布的系列传播内容", "短视频8—12期", "内容台账、发布记录"),
        ],
        [0.9, 2.25, 1.35, 1.5],
        font_size=8.7,
    )
    add_callout(doc, "口径说明｜", "上述为项目目标值，不得在申报书中表述为已取得成果；申报前30—50人测试用于校准指标。")

    add_heading(doc, "2. 用户、场景与产品原则", 1)
    add_heading(doc, "2.1 核心用户与使用场景", 2)
    add_labeled_paragraph(doc, "核心用户｜", "广东高校学生，重点为18—24岁青年；产品不设置年龄排他限制。")
    add_labeled_paragraph(doc, "典型场景｜", "在短视频、海报、课程或校园活动中扫码进入；利用课间或碎片时间完成体验；体验后自主选择继续了解或参加活动。")
    add_labeled_paragraph(doc, "用户特征｜", "对新奇互动有兴趣，但对长篇学习和强制答题耐心有限；愿意分享个性化结果，但不希望先完成复杂注册。")

    add_heading(doc, "2.2 产品原则", 2)
    for item in [
        "先体验后学习：首屏直接提供“开始3分钟体验”，知识内容作为体验后的自愿延伸。",
        "真实内容优先：工艺、人物和史实必须有来源并经专业审核，数字互动明确标注为简化模拟。",
        "引导完成优先：核心体验不设失败、倒计时、扣分和数值排名，允许提示与重试。",
        "核心流程可降级：3D加载失败或设备能力不足时提供2D、图文或视频替代路径。",
        "最少干扰：核心流程无需注册；相机权限仅在用户点击AR时请求。",
    ]:
        add_bullet(doc, item, principles_bullets)

    add_heading(doc, "3. 范围与版本优先级", 1)
    add_table(
        doc,
        ["级别", "版本目标", "纳入功能", "完成定义"],
        [
            ("P0", "申报演示", "首页入口、分段真实视频、3步3D互动、即时反馈、结果页、3D降级、核心埋点", "可在真实手机完成主流程并形成演示数据"),
            ("P1", "正式首版", "传承人故事、金箔志、一个AR场景、活动、分享海报、内容导入发布", "内容经审核、兼容测试通过、可开展校园验证"),
            ("P2", "后续迭代", "登录、积分徽章、排行、复杂测评、收藏、多非遗、商城支付、直播", "不进入首版工期与验收"),
        ],
        [0.65, 1.1, 2.9, 1.35],
        font_size=8.6,
    )
    add_callout(doc, "范围冻结｜", "P0验收前不得启动P2。P1新增功能须先证明不会影响核心体验、兼容性与用户研究进度。", fill=CREAM)

    add_heading(doc, "4. 信息架构与核心流程", 1)
    add_heading(doc, "4.1 一级导航", 2)
    add_table(
        doc,
        ["入口", "主要任务", "核心内容"],
        [
            ("首页", "快速理解价值并开始体验", "主视觉、3分钟承诺、开始按钮、精选内容与活动入口"),
            ("体验", "完成3步轻量工艺互动", "分段视频、3D操作、即时反馈、结果与延伸动作"),
            ("金箔志", "按兴趣了解真实人物与工艺", "人物、工艺、工具与材料、作品与应用"),
            ("活动", "了解并进入线下参与渠道", "活动列表、详情、官方报名二维码或渠道"),
        ],
        [0.9, 2.2, 2.9],
        font_size=9.0,
    )
    add_body(doc, "首版不设置“我的”一级入口。登录、收藏、积分、证书和用户资产均不属于P0/P1必需能力。")

    add_heading(doc, "4.2 核心用户流程", 2)
    for item in [
        "用户从首页点击“开始3分钟体验”，无需注册即可进入。",
        "观看3—5秒真实人物开场；可跳过，字幕默认可读，静音不影响理解。",
        "在每个互动步骤前观看5—8秒真实工艺片段，再完成对应3D操作。",
        "系统提供即时描述性反馈、提示和重试，完成3步后生成个人金箔作品。",
        "结果页提供AR、分享海报、金箔志、活动和再次体验等可选动作。",
    ]:
        add_numbered(doc, item, flow_numbers)
    add_callout(doc, "核心闭环｜", "首页 → 真实开场 → 视频片段＋3D步骤×3 → 个人作品与反馈 → 可选深度内容／AR／活动。")

    add_heading(doc, "4.3 工艺主线决策框架", 2)
    add_body(doc, "在首次访谈前，产品仅保留共享的三阶段交互框架。访谈确认后，将下表中一列锁定为首版内容，并由专业老师审核术语、步骤与表现。")
    add_table(
        doc,
        ["共享阶段", "候选A：制箔", "候选B：贴箔", "共享产品能力"],
        [
            ("1 认识任务", "认识材料、工具与目标厚度", "选择器物、金箔与工具", "真实短视频、3D对象、提示卡"),
            ("2 关键动作", "模拟敲击力度与节奏", "模拟取箔、移动与覆盖", "手势识别、即时反馈、重试"),
            ("3 检查完成", "观察厚度、完整性与裁切", "刷平、修补并检查覆盖", "完成状态、个性结果、文化提示"),
        ],
        [1.05, 1.55, 1.55, 1.85],
        font_size=8.6,
    )
    add_labeled_paragraph(doc, "锁定依据｜", "合作对象的代表性工艺、可拍摄性、3分钟内的可交互性、文化准确性和技术可实现性。")

    add_heading(doc, "5. 功能需求", 1)
    add_heading(doc, "5.1 首页", 2)
    add_requirement_table(doc, [
        ("HOME-01", "首屏展示一句话价值和“开始3分钟体验”主按钮。", "首次进入无需滚动即可看到并点击主按钮。"),
        ("HOME-02", "首屏不以课程、知识测评或注册作为前置。", "未登录用户可直接进入EXP-01。"),
        ("HOME-03", "展示精选故事、金箔志和近期活动入口。", "内容为空时模块可隐藏，不影响主按钮。"),
        ("HOME-04", "记录主按钮点击来源。", "触发home_cta_click并携带入口来源。"),
    ])

    add_heading(doc, "5.2 分段视频与体验准备", 2)
    add_requirement_table(doc, [
        ("EXP-01", "进入体验后播放3—5秒传承人或专业老师开场。", "视频可跳过，有字幕，静音状态可理解。"),
        ("EXP-02", "每个步骤前播放5—8秒真实工艺片段。", "每段可跳过；结束或跳过后进入对应操作。"),
        ("EXP-03", "首次加载显示清晰进度与预计等待说明。", "用户不会停留在无反馈空白页。"),
        ("EXP-04", "说明体验为简化数字模拟。", "开始前或首步可见，不宣称替代真实工艺。"),
    ])

    add_heading(doc, "5.3 三步3D互动", 2)
    add_requirement_table(doc, [
        ("INT-01", "按已锁定工艺主线提供连续3步操作。", "步骤顺序、术语和动作经专业审核。"),
        ("INT-02", "每步提供首次引导和按需提示。", "无操作或多次偏离时出现提示；用户可继续。"),
        ("INT-03", "操作后立即给出视觉、文字或动效反馈。", "不以声音或颜色作为唯一反馈。"),
        ("INT-04", "不设置倒计时、失败、扣分或公开排名。", "用户通过提示和重试可完成全部步骤。"),
        ("INT-05", "后台记录步骤时长、提示次数和重试次数。", "不向用户展示数值分数，但事件字段可查询。"),
        ("INT-06", "支持退出并记录退出步骤。", "退出不进入死循环；触发experience_exit。"),
    ], font_size=8.6)

    add_heading(doc, "5.4 结果页", 2)
    add_requirement_table(doc, [
        ("RES-01", "生成并展示个人3D金箔作品。", "完成3步后自动进入结果页并可旋转查看。"),
        ("RES-02", "根据操作过程给出描述性反馈。", "反馈使用“细致、稳定、均匀”等措辞，不显示分数或排名。"),
        ("RES-03", "展示一条经审核的文化知识提示。", "知识内容可追溯至内容台账和审核记录。"),
        ("RES-04", "提供AR、分享、金箔志、活动和重试入口。", "各动作可选；拒绝权限或分享失败不影响结果。"),
        ("RES-05", "分享海报不得强制填写姓名、手机号。", "可匿名生成；分享动作由用户主动发起。"),
    ], font_size=8.6)

    add_heading(doc, "5.5 金箔志", 2)
    add_requirement_table(doc, [
        ("ARC-01", "设置人物、工艺、工具与材料、作品与应用四类内容。", "分类入口均可独立配置和隐藏空分类。"),
        ("ARC-02", "列表卡片展示标题、摘要、封面与类型。", "点击进入1分钟可读详情。"),
        ("ARC-03", "详情支持短文、图片与可选长视频。", "视频非自动播放；弱网下正文仍可阅读。"),
        ("ARC-04", "详情底部显示来源、审核人和审核日期。", "未审核内容不得进入已发布状态。"),
        ("ARC-05", "P1仅提供分类筛选，不做全文搜索。", "四类筛选结果正确；无搜索入口。"),
    ], font_size=8.5)

    add_heading(doc, "5.6 活动、AR与分享", 2)
    add_requirement_table(doc, [
        ("ACT-01", "活动模块展示列表、详情、时间地点和主办方。", "信息与官方发布一致，过期状态明确。"),
        ("ACT-02", "点击“我要参加”后展示学校或主办方官方报名渠道。", "小程序仅记录意向点击，不代收报名信息。"),
        ("AR-01", "结果页提供一个可选AR展示场景。", "仅点击后申请相机权限；拒绝后返回结果页。"),
        ("SHR-01", "生成包含作品、项目识别和小程序入口的分享海报。", "分享失败可重试或返回，不阻断其他动作。"),
    ], font_size=8.5)

    add_heading(doc, "6. 状态、异常与降级", 1)
    add_table(
        doc,
        ["场景", "系统行为", "验收要求"],
        [
            ("3D加载中", "显示进度和操作提示", "无空白页；用户可返回"),
            ("加载超过8秒", "提示继续等待并说明可退出", "不重复弹窗、不误判完成"),
            ("加载15秒失败", "提供重试或2D／图文／视频替代", "替代路径可完成基础内容体验"),
            ("弱网或断网", "提示网络状态并允许重试", "恢复网络后无需重启小程序"),
            ("拒绝相机权限", "关闭AR并返回结果页", "不再次强制请求，不影响其他功能"),
            ("视频加载失败", "显示封面、字幕摘要与跳过", "用户仍可进入对应3D步骤"),
            ("内容为空或下架", "隐藏入口或显示友好空状态", "不出现无效链接和报错代码"),
        ],
        [1.25, 2.8, 1.95],
        font_size=8.8,
    )

    add_heading(doc, "7. 数据埋点与效果评价", 1)
    add_heading(doc, "7.1 核心事件", 2)
    add_table(
        doc,
        ["事件", "触发时机", "关键字段", "用途"],
        [
            ("home_cta_click", "点击首页主按钮", "source", "入口转化"),
            ("experience_start", "进入核心体验", "session_id、route", "开始人数"),
            ("video_play / skip", "播放或跳过分段视频", "segment_id、watch_time", "内容接受度"),
            ("step_start", "进入某一步", "step_id", "漏斗起点"),
            ("step_hint / retry", "使用提示或重试", "step_id、count", "难度诊断"),
            ("step_complete", "完成某一步", "step_id、duration", "步骤完成率"),
            ("experience_exit", "中途退出", "step_id、duration", "流失诊断"),
            ("experience_complete", "完成全部3步", "total_duration", "核心完成率"),
            ("result_action", "结果页点击后续动作", "action_type", "深度探索"),
            ("archive_view", "打开金箔志详情", "content_id、category", "内容偏好"),
            ("activity_interest", "点击我要参加", "activity_id", "活动意向"),
        ],
        [1.35, 1.65, 1.65, 1.35],
        font_size=8.1,
    )
    add_labeled_paragraph(doc, "完成率｜", "experience_complete 去重会话数 ÷ experience_start 去重会话数。")
    add_labeled_paragraph(doc, "继续探索率｜", "完成体验后触发AR、分享、金箔志或活动动作的去重用户数 ÷ 完成体验去重用户数。")
    add_labeled_paragraph(doc, "数据质量｜", "事件均需包含时间、版本、设备能力、会话编号和匿名用户编号；测试账号与正式数据应可区分。")

    add_heading(doc, "7.2 两层用户研究设计", 2)
    add_table(
        doc,
        ["层级", "对象与规模", "收集内容", "用途"],
        [
            ("一般用户层", "约800名真实用户", "匿名行为；体验后自愿填写学校、年级、专业大类、兴趣", "评价触达、完成率和探索行为"),
            ("正式研究层", "不少于100名知情同意样本", "手机号或微信联系方式、前测、后测、随访、行为事件", "配对分析、访谈和论文研究"),
        ],
        [1.1, 1.55, 2.2, 1.15],
        font_size=8.7,
    )

    add_heading(doc, "8. 隐私、同意与数据分离", 1)
    for item in [
        "核心体验不以填写手机号、微信号或研究问卷为前提；非必要信息由用户自愿提供。",
        "正式研究样本在收集联系方式和配对数据前展示研究目的、使用范围、保存期限、退出方式和联系人，并取得明确同意。",
        "活动报名继续使用学校或主办方官方渠道；若活动参与者同时进入研究，需另行完成研究同意。",
        "联系方式仅用于约定的随访、访谈或通知，不得转作营销、公开传播或与无关数据合并。",
        "原始研究数据仅限授权成员访问；对外报告和论文使用去标识化或汇总结果。",
    ]:
        add_bullet(doc, item, privacy_bullets)
    add_table(
        doc,
        ["数据表", "建议字段", "访问要求"],
        [
            ("research_contacts", "research_id、手机号／微信、同意时间、联系状态", "单独存储，仅研究负责人和指导教师授权访问"),
            ("research_responses", "research_id、前测、后测、随访答案", "不存直接联系方式；按研究编号关联"),
            ("behavior_events", "anonymous_id／research_id、事件、时间、版本、参数", "研究样本经同意后才关联research_id"),
        ],
        [1.35, 3.0, 1.65],
        font_size=8.7,
    )
    add_callout(doc, "合规基线｜", "目的明确、范围合理、最少必要、知情自愿、明确同意、可撤回、限制访问。是否愿意留下手机号，不替代告知与同意程序。", fill=CREAM)

    add_heading(doc, "9. 内容生产、审核与发布", 1)
    add_heading(doc, "9.1 首版工作流", 2)
    for item in [
        "岗位3在结构化Excel／CSV内容台账中建档，分配内容ID和素材ID。",
        "传承人或专业老师审核工艺、人物与史实；团队核对授权、来源和文字规范。",
        "内容状态依次流转为草稿、待审核、已审核、已发布；只有已审核内容可导入。",
        "岗位1或岗位2使用受控导入脚本写入云数据库，媒体文件上传云存储。",
        "发布后保留版本、修改人、修改时间、修改原因和下架记录。",
    ]:
        add_numbered(doc, item, content_numbers)
    add_callout(doc, "首版约束｜", "不建设复杂CMS，也不允许内容学生直接在云数据库中随意修改。云控制台仅用于技术维护和紧急处置。")

    add_heading(doc, "9.2 内容台账最低字段", 2)
    add_table(
        doc,
        ["字段组", "最低字段", "规则"],
        [
            ("识别", "content_id、category、title、summary", "ID唯一且发布后不复用"),
            ("正文", "body、media_ids、related_ids", "正文引用素材ID，不直接写本地路径"),
            ("真实性", "source、reviewer、review_date", "来源可追溯；审核人和日期不可为空"),
            ("授权", "authorization_id、rights_scope", "图片、视频、肖像、模型均需对应授权或许可"),
            ("状态", "status、version、updated_by、updated_at", "仅reviewed可导入published"),
        ],
        [1.0, 3.2, 1.8],
        font_size=8.6,
    )

    add_heading(doc, "10. 技术方案与数据结构", 1)
    add_heading(doc, "10.1 技术架构", 2)
    add_table(
        doc,
        ["层级", "选型", "首版职责"],
        [
            ("客户端", "微信原生小程序", "页面、状态管理、视频、分享、权限与降级"),
            ("3D／AR", "微信XR-Frame＋优化后的GLB模型", "3步互动、结果作品、一个可选AR场景"),
            ("云服务", "腾讯云开发CloudBase", "云数据库、云存储、云函数、内容读取与数据写入"),
            ("内容导入", "Excel／CSV＋受控导入脚本", "审核后批量校验、导入、版本记录"),
            ("媒体策略", "云存储＋按需加载／分包", "视频、图片和3D资源不阻塞首页"),
        ],
        [1.0, 2.1, 2.9],
        font_size=8.8,
    )
    add_body(doc, "P0/P1不引入PHP、MySQL、自建服务器或uni-app，不建设复杂后台管理系统。若后续内容量和运营频率显著增加，再单独评估轻量管理端。")

    add_heading(doc, "10.2 建议集合", 2)
    add_table(
        doc,
        ["集合", "用途", "责任岗位"],
        [
            ("app_config", "版本、功能开关、首页推荐和工艺路线配置", "岗位1／2"),
            ("content_items", "人物、工艺、工具材料、作品应用内容", "岗位3提供，岗位1／2导入"),
            ("media_assets", "图片、视频、模型及授权引用", "岗位2／3"),
            ("activities", "活动信息、状态和官方报名渠道", "岗位4提供，岗位1导入"),
            ("behavior_events", "匿名或研究编号关联的行为事件", "岗位2／4"),
            ("research_contacts", "正式研究样本联系方式与同意记录", "岗位4，限制访问"),
            ("research_responses", "前后测、随访和访谈结构化数据", "岗位4"),
            ("audit_records", "内容版本、导入、修改和下架记录", "岗位2／3"),
        ],
        [1.4, 3.2, 1.4],
        font_size=8.4,
    )
    add_callout(doc, "设计边界｜", "本节只定义逻辑集合。字段类型、索引、权限规则和云函数接口应在《技术设计说明书》中进一步细化。", fill=SOFT_GRAY)

    add_heading(doc, "11. 非功能需求", 1)
    add_table(
        doc,
        ["类别", "内部目标", "验收方式"],
        [
            ("首页性能", "校园Wi-Fi／4G下2.5秒内可用", "6台真机冷启动与二次启动记录"),
            ("3D性能", "中档手机6秒内可操作；全程显示进度", "记录加载时长、帧率和失败原因"),
            ("资源体积", "核心GLB目标≤6MB；单场景网络资源目标≤12MB", "构建产物和网络面板检查"),
            ("视频", "步骤视频5—8秒、有字幕、静音可理解", "内容检查与用户测试"),
            ("兼容", "iOS／Android高、中、低档共6台真机", "设备矩阵、截图和问题闭环"),
            ("可访问性", "触控区域清晰、对比度足够、不依赖单一颜色或声音", "页面与交互走查"),
            ("稳定性", "核心流程无崩溃、死循环和不可退出状态", "弱网、断网、权限拒绝与异常测试"),
            ("安全与权限", "相机仅按需申请；敏感集合最小权限", "权限规则与访问测试"),
        ],
        [1.0, 3.05, 1.95],
        font_size=8.5,
    )
    add_callout(doc, "校准规则｜", "以上为内部工程目标，完成技术原型后根据真实设备和资源复杂度校准，但不得取消降级、权限和稳定性要求。")

    add_heading(doc, "12. 验收与测试", 1)
    add_heading(doc, "12.1 P0申报演示验收", 2)
    for item in [
        "从首页无需注册进入体验，完整完成真实开场、3段短视频、3步3D操作和结果页。",
        "3D加载失败、视频失败和弱网情况下均有可理解提示、重试或替代内容。",
        "核心事件在测试环境可查询，完成率漏斗与退出步骤可计算。",
        "工艺路线、术语、视频和文化提示已取得专业审核记录。",
        "至少完成一轮30—50名高校学生测试并形成问题、改进和数据摘要。",
    ]:
        add_bullet(doc, item, testing_bullets)

    add_heading(doc, "12.2 P1正式首版验收矩阵", 2)
    add_table(
        doc,
        ["验收域", "必须通过的检查", "主要责任"],
        [
            ("功能", "HOME、EXP、INT、RES、ARC、ACT、AR、SHR需求逐项通过", "岗位1／2"),
            ("内容", "已发布内容均有来源、授权、审核人、日期和版本", "岗位3"),
            ("设备", "6台真机覆盖微信版本、系统、性能和相机权限", "岗位2"),
            ("数据", "事件不漏报、不重复；测试数据可识别；指标可复算", "岗位2／4"),
            ("研究", "同意文本、研究编号、三表分离和退出机制可用", "岗位4"),
            ("实践", "活动渠道有效，意向点击与主办方汇总报名数可对照", "岗位4"),
        ],
        [1.0, 3.7, 1.3],
        font_size=8.6,
    )

    add_heading(doc, "13. 四岗位责任与交接", 1)
    add_table(
        doc,
        ["岗位", "在本PRD中的主责", "必须交接的材料"],
        [
            ("1 产品与前端", "首页、导航、页面流程、视频、结果、金箔志、活动和前端埋点", "原型、页面清单、组件说明、验收截图、使用说明"),
            ("2 3D与技术", "XR-Frame、3D交互、AR、CloudBase、导入、权限、性能和数据", "模型规范、接口、集合、部署、设备测试和故障说明"),
            ("3 内容与档案", "访谈、拍摄、内容台账、授权、专业审核和发布素材", "内容手册、素材ID、来源授权、审核与版本记录"),
            ("4 调研与实践", "研究设计、同意、问卷、活动、数据分析和证据归档", "研究方案、数据字典、活动材料、分析报告和证据目录"),
        ],
        [1.15, 3.0, 1.85],
        font_size=8.5,
    )
    add_labeled_paragraph(doc, "共同责任｜", "四名学生共同参加首次专业访谈、P0评审、用户测试复盘和P1验收；主责人对结果负责，备份人能够接手。")

    add_heading(doc, "14. 待锁定事项与下一步", 1)
    add_table(
        doc,
        ["事项", "锁定时间", "责任人", "所需证据"],
        [
            ("首版采用制箔或贴箔主线", "首次访谈后48小时内", "指导教师＋岗位2／3", "访谈纪要、专业确认、技术评估"),
            ("非遗项目正式名称与认定信息", "申报准备第1—2周", "岗位3", "官方名录或保护单位材料"),
            ("小程序正式名称与视觉识别", "低保真原型评审时", "岗位1／3", "命名清单、合作方意见"),
            ("AR展示对象与最低设备基线", "3D技术原型评审时", "岗位2", "模型样例、6台设备测试"),
            ("正式研究问卷与同意文本", "首轮正式研究前", "岗位4＋指导教师", "预测试、修改记录、定稿文件"),
        ],
        [2.0, 1.4, 1.2, 1.4],
        font_size=8.5,
    )
    add_body(doc, "工艺路线锁定后，团队应立即完成交互脚本、素材清单、3D资产清单和P0任务拆解，并将本PRD版本更新为V0.2。")

    add_heading(doc, "建议执行顺序", 2)
    for item in [
        "完成首次访谈并锁定制箔或贴箔主线。",
        "岗位1产出低保真流程原型，岗位3同步完成对应内容与拍摄清单。",
        "岗位2制作一个步骤的3D技术原型并在6台设备中先测中低档机。",
        "岗位4定义埋点字段、前后测和测试招募流程。",
        "四岗位共同评审后进入P0开发，按需求编号逐项验收。",
    ]:
        add_numbered(doc, item, decision_numbers)

    add_heading(doc, "附录：技术与合规参考", 1)
    add_body(doc, "微信XR-Frame官方文档：https://developers.weixin.qq.com/miniprogram/dev/component/xr-frame/")
    add_body(doc, "腾讯云开发CloudBase官方文档：https://cloud.tencent.cn/document/product/876/121103")
    add_body(doc, "《中华人民共和国个人信息保护法》：https://www.npc.gov.cn/npc/c2/c30834/202108/t20210820_313088.html")
    add_callout(doc, "版本说明｜", "PRD V0.1固化截至当前已确认的产品决策；涉及合作对象身份、正式工艺路线和真实素材的内容，须以首次专业访谈及书面审核为准。", fill=CREAM)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
