from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_gold_foil_project_brief_docx import (  # noqa: E402
    CREAM,
    CREAM_LIGHT,
    GOLD,
    GOLD_DARK,
    GRID,
    INK,
    MUTED,
    SOFT_GRAY,
    WHITE,
    add_body,
    add_callout,
    add_field,
    fill_cell,
    prevent_row_split,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_style_font,
    set_table_borders,
)

TABLE_GEOMETRY_DIR = Path(
    r"C:\Users\Andrew\.codex\plugins\cache\openai-primary-runtime\documents\26.715.12143\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_GEOMETRY_DIR))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


OUT = ROOT / "docs" / "gold-foil-heritage" / "金箔非遗项目四岗位7周任务与验收表V0.1.docx"


WEEKS = [
    {
        "week": "第1周",
        "dates": "7月17日—7月23日",
        "goal": "完成专业访谈并锁定工艺主线，同时建立授权、技术和研究基线。",
        "rows": [
            (
                "1 产品与前端",
                "参加专业访谈，整理目标用户、核心流程、P0页面和待决策项；建立产品决策记录。",
                "依赖岗位3的访谈安排与纪要；与岗位2共同判断交互可实现性。",
                "P0页面清单V0.1、待决策清单、产品决策记录。",
                "覆盖首页、体验、结果页及异常入口；每个待决策项均有责任人、期限和证据要求。",
                "指导教师＋岗位2",
                "未开始",
            ),
            (
                "2 3D与技术",
                "建立微信小程序、CloudBase和XR-Frame基础环境；准备6台真机清单；评估制箔与贴箔的技术差异。",
                "依赖岗位1的核心流程；从岗位3获得候选动作、工具和材料信息。",
                "可运行工程骨架、XR基础场景、设备清单、技术可行性记录。",
                "至少1台真机可运行XR基础场景；明确两条候选路线的模型、手势、性能和降级风险。",
                "指导教师＋岗位1",
                "未开始",
            ),
            (
                "3 内容与档案",
                "组织首次访谈，核实项目和合作对象信息，明确可拍摄工艺、公开边界与授权范围。",
                "全体参加关键访谈；岗位2提供技术问题，岗位4提供研究问题。",
                "访谈纪要、专业确认、身份依据清单、授权事项清单、首批素材目录。",
                "制箔或贴箔主线有专业依据；每类素材标明可拍、可公开、需补授权或不可使用。",
                "指导教师",
                "未开始",
            ),
            (
                "4 调研与实践",
                "定义一般用户层与正式研究层的目标、样本口径、招募路径和前后测构念。",
                "依赖岗位1的用户路径；与岗位3确认访谈和活动可接触人群。",
                "研究目标说明、样本口径表、招募计划、前后测构念清单。",
                "明确800用户、100配对样本和30—50人测试的定义、排除规则与证据来源。",
                "指导教师",
                "未开始",
            ),
        ],
    },
    {
        "week": "第2周",
        "dates": "7月24日—7月30日",
        "goal": "完成低保真方案、拍摄与资产清单，启动单步骤3D原型和研究工具设计。",
        "rows": [
            (
                "1 产品与前端",
                "制作首页至结果页的低保真可点击原型，补齐加载、空状态、退出、权限拒绝和降级路径。",
                "工艺主线已锁定；岗位2反馈技术边界，岗位3审核文化表述。",
                "低保真原型V0.1、页面状态清单、评审问题单。",
                "主流程无需注册可完成；覆盖HOME、EXP、INT、RES及主要异常状态；评审问题均有结论。",
                "指导教师＋岗位2／3",
                "未开始",
            ),
            (
                "2 3D与技术",
                "制定模型、纹理、动画、手势和资源体积规范；启动一个代表性步骤的技术原型。",
                "依赖岗位3提供工具、材料和动作参考；与岗位1确认反馈方式。",
                "3D资产规范V0.1、GLB资产清单、单步骤原型工程。",
                "核心GLB目标不超过6MB；交互动作、反馈、提示和降级方案均写入规范。",
                "岗位1＋指导教师",
                "未开始",
            ),
            (
                "3 内容与档案",
                "完成开场及三步骤拍摄脚本、镜头清单、采访提纲和内容台账；推进书面授权。",
                "依赖岗位1的低保真流程和岗位2的动作／资产需求。",
                "拍摄执行单、采访提纲、内容台账V0.1、授权文件清单。",
                "每个视频段落对应一个体验节点；素材均有内容ID、来源、授权状态和审核人字段。",
                "指导教师＋岗位1／2",
                "未开始",
            ),
            (
                "4 调研与实践",
                "形成埋点字典和问卷初稿，定义渠道来源、完成率、继续探索率及研究编号规则。",
                "依赖岗位1页面清单和岗位2事件实现方式；与岗位3核对知识题。",
                "埋点字典V0.1、问卷V0.1、知情同意文本V0.1、招募话术。",
                "事件覆盖home_cta_click至activity_interest；联系方式、问卷和行为数据分表设计明确。",
                "指导教师＋岗位2",
                "未开始",
            ),
        ],
    },
    {
        "week": "第3周",
        "dates": "7月31日—8月6日",
        "goal": "完成单步骤技术验证、正式拍摄、前端骨架和研究预测试。",
        "rows": [
            (
                "1 产品与前端",
                "建立首页、体验、结果页和导航骨架；确定基础视觉组件和前端埋点接入位置。",
                "使用已评审低保真原型；与岗位2约定页面和3D场景接口。",
                "前端工程骨架、组件清单、可点击高保真关键页面、接口约定。",
                "首页可进入体验占位页；返回、退出和降级入口可用；组件命名和状态统一。",
                "岗位2＋指导教师",
                "未开始",
            ),
            (
                "2 3D与技术",
                "完成单步骤3D原型、CloudBase基础集合和6台设备中的中低档机优先测试。",
                "依赖岗位3的动作参考与岗位1的交互反馈；向岗位4提供事件字段。",
                "单步骤可操作原型、设备测试记录、性能数据、技术路线结论、埋点字典技术稿。",
                "中档机6秒目标有实测记录；失败时有点击／滑动或2D替代；主线技术路线可进入P0开发。",
                "指导教师＋岗位1",
                "未开始",
            ),
            (
                "3 内容与档案",
                "实施正式拍摄并完成双份备份；整理开场和三步骤素材，剪出一个示例片段。",
                "依赖拍摄授权和技术镜头清单；岗位4协助现场记录。",
                "原始素材包、备份清单、素材ID目录、示例短片、拍摄现场记录。",
                "关键镜头完整、声音可用、人物和工艺信息一致；所有素材可追溯至授权与来源。",
                "指导教师＋岗位1／2",
                "未开始",
            ),
            (
                "4 调研与实践",
                "对问卷和知情同意文本进行8—12人预测试；冻结渠道枚举和数据质量规则。",
                "依赖岗位2的事件字段和岗位3的知识内容；测试对象不计入正式样本。",
                "预测试记录、问卷V0.2、同意文本V0.2、埋点字典定稿。",
                "题目无明显歧义；前后测可配对；渠道、测试账号、无效数据和退出样本规则明确。",
                "指导教师",
                "未开始",
            ),
        ],
    },
    {
        "week": "第4周",
        "dates": "8月7日—8月13日",
        "goal": "完成P0三步流程、视频成片与审核，并冻结测试工具。",
        "rows": [
            (
                "1 产品与前端",
                "开发首页、分段视频、结果页和导航；接入3D场景、内容配置、权限与异常提示。",
                "依赖岗位2提供三步场景接口，岗位3提供审核视频和文案。",
                "HOME／EXP／RES功能、前端埋点、分享与降级入口。",
                "HOME-01至HOME-04、EXP-01至EXP-04、RES-01至RES-05可按验收口径逐项检查。",
                "指导教师＋岗位2",
                "未开始",
            ),
            (
                "2 3D与技术",
                "实现三步完整3D流程、提示重试、描述性反馈、结果作品、核心埋点和加载降级。",
                "依赖岗位1的页面接口和岗位3的专业动作审核；与岗位4核对数据。",
                "INT三步场景、结果作品、行为事件、加载和2D／视频降级逻辑。",
                "INT-01至INT-06通过；三步均可完成，至少1步高保真；简化步骤不得取消真实视频、即时反馈和结果页。",
                "指导教师＋岗位1／4",
                "未开始",
            ),
            (
                "3 内容与档案",
                "完成开场和三段步骤视频剪辑、字幕、解说、专业审核及小程序导入包。",
                "岗位1提供画幅与时长规范；岗位2核对动作节奏；专业老师完成审核。",
                "4段审核视频、字幕文件、审核记录、内容台账V0.2、导入素材包。",
                "开场3—5秒、步骤5—8秒；静音可理解；每项内容有来源、授权、审核人和日期。",
                "指导教师＋专业老师",
                "未开始",
            ),
            (
                "4 调研与实践",
                "定稿问卷、同意文本、测试方案和招募名单；建立匿名编号与研究编号管理表。",
                "依赖岗位1／2提供可测试流程和事件版本；岗位3审核知识题。",
                "问卷定稿、同意文本定稿、测试主持稿、招募与编号规则。",
                "一般用户不强制留联系方式；正式研究样本可撤回；三类数据可按研究编号安全关联。",
                "指导教师",
                "未开始",
            ),
        ],
    },
    {
        "week": "第5周",
        "dates": "8月14日—8月20日",
        "goal": "完成P0联调、体验版发布、6台真机测试和首轮30—50人用户测试。",
        "rows": [
            (
                "1 产品与前端",
                "完成前端联调和体验版发布，修复主流程、文案、页面跳转和用户可理解性问题。",
                "依赖岗位2部署和岗位3内容导入；支持岗位4现场测试。",
                "体验版二维码、前端问题清单、修复记录、关键页面截图。",
                "从首页到结果页无需注册可完整体验；空状态、返回和分享失败不阻断主流程。",
                "指导教师＋岗位4",
                "未开始",
            ),
            (
                "2 3D与技术",
                "开展6台真机、弱网、断网、相机拒绝、视频失败和埋点准确性测试。",
                "与岗位1联合联调；向岗位4提供可查询事件和测试账号标记。",
                "设备矩阵、性能记录、异常测试、事件核验表、部署版本记录。",
                "无崩溃和死循环；15秒失败可降级；事件不漏报、不重复，测试数据可识别。",
                "指导教师＋岗位1／4",
                "未开始",
            ),
            (
                "3 内容与档案",
                "完成上线内容、字幕、授权和文化准确性终检；现场观察用户对视频和术语的理解。",
                "依赖岗位1展示效果和岗位4测试安排；问题回传给专业老师。",
                "上线内容清单、内容终检表、授权索引、用户内容反馈记录。",
                "公开内容全部处于已审核状态；未授权或待确认内容不进入体验版；反馈问题可追踪。",
                "指导教师＋岗位4",
                "未开始",
            ),
            (
                "4 调研与实践",
                "组织30—50名高校学生测试，记录完成、退出、继续探索、问卷与访谈反馈。",
                "岗位1／2提供稳定体验版，岗位3协助现场说明；研究补偿与传播激励分开。",
                "招募记录、匿名行为数据、问卷数据、访谈纪要、现场照片和问题优先级。",
                "样本来源与排除规则可追溯；联系方式仅在明确同意的正式研究中收集；问题按严重度分级。",
                "指导教师",
                "未开始",
            ),
        ],
    },
    {
        "week": "第6周",
        "dates": "8月21日—8月27日",
        "goal": "修复高优先级问题，完成性能降级、内容收尾和申报数据摘要。",
        "rows": [
            (
                "1 产品与前端",
                "修复P0高优先级问题并冻结页面结构；完成演示路径、使用说明和申报截图。",
                "依赖岗位4的问题优先级和岗位2的版本；岗位3复核修改文案。",
                "P0候选版本、修复清单、演示路径、使用说明初稿、截图包。",
                "阻断和高优先级问题关闭；页面与PRD一致；新增功能未突破P0范围。",
                "指导教师＋岗位2／4",
                "未开始",
            ),
            (
                "2 3D与技术",
                "压缩模型和纹理，优化加载与帧率，完善视频／2D降级、日志、部署和备份。",
                "依据第5周设备记录；与岗位1共同验证降级体验。",
                "优化版本、资源体积表、降级测试记录、部署与恢复说明。",
                "核心GLB和场景资源达到校准目标；中低档机可完成；部署失败可恢复至上一版本。",
                "指导教师＋岗位1",
                "未开始",
            ),
            (
                "3 内容与档案",
                "完成内容台账、素材ID、审核、授权和版本记录收尾，整理申报用人物与工艺证据。",
                "根据测试反馈修改内容；与岗位1确认最终展示版本。",
                "内容证据包、授权索引、审核记录、最终素材目录、精选申报素材。",
                "已发布内容审核状态清零；每项申报素材可追溯；不得使用未确认称号和级别。",
                "指导教师",
                "未开始",
            ),
            (
                "4 调研与实践",
                "清洗测试数据并计算漏斗、退出步骤和继续探索率，形成问题与效果摘要。",
                "依赖岗位2事件数据和岗位1版本号；与岗位3解释内容反馈。",
                "数据字典、清洗记录、漏斗表、用户反馈摘要、改进闭环表。",
                "指标可按原始数据复算；测试账号和无效样本已排除；目标值与既有成果表述严格区分。",
                "指导教师＋岗位2",
                "未开始",
            ),
        ],
    },
    {
        "week": "第7周",
        "dates": "8月28日—9月3日",
        "goal": "冻结P0版本，完成部署备份、申报材料、演示脚本和答辩演练。",
        "rows": [
            (
                "1 产品与前端",
                "冻结P0版本，完成演示操作脚本、关键截图、用户说明和现场演示分工。",
                "与岗位2确认部署版本；使用岗位3／4提供的内容和证据。",
                "P0正式候选版、演示脚本、截图包、使用说明、版本说明。",
                "12.1节P0验收全部通过；3分钟主流程可连续演示；现场操作不依赖临时修改。",
                "指导教师",
                "未开始",
            ),
            (
                "2 3D与技术",
                "完成部署、云端和本地备份、离线演示预案、故障恢复与技术答辩材料。",
                "与岗位1联合演练；从岗位4获取数据口径。",
                "部署记录、备份包、离线演示视频、故障清单、技术答辩卡。",
                "网络异常仍有可展示方案；备份可恢复；技术指标与实测记录一致，不夸大能力。",
                "指导教师＋岗位1",
                "未开始",
            ),
            (
                "3 内容与档案",
                "整理传承人故事、工艺依据、授权和内容审核材料，完成答辩用视觉与视频素材。",
                "与岗位1确认演示画面；向岗位4提供社会文化价值证据。",
                "内容答辩包、精选视频、人物与工艺证据、授权和审核索引。",
                "所有身份、级别和工艺表述有依据；视频与图片授权明确；故事与产品体验一致。",
                "指导教师＋专业老师",
                "未开始",
            ),
            (
                "4 调研与实践",
                "完成申报数据、获客路径、风险状态、PPT证据页和答辩问答，组织全流程演练。",
                "整合岗位1／2／3交付；指导教师校准表述。",
                "数据摘要、风险状态表、获客计划、PPT证据页、答辩题库和演练记录。",
                "目标、测试结果和未来计划严格区分；核心数据可追溯；至少完成2轮计时演练。",
                "指导教师",
                "未开始",
            ),
        ],
    },
]


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    set_style_font(normal, 9.5, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.18

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 16, INK, bold=True)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 12.5, GOLD_DARK, bold=True)
    h2.paragraph_format.space_before = Pt(7)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("金箔非遗数字体验与活化传播项目"), size=8.2, bold=True, color=MUTED)
    set_run_font(header.add_run("    |    四岗位7周任务与验收表 V0.1"), size=8.2, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("第 "), size=8.5, color=MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.add_run(" 页"), size=8.5, color=MUTED)

    doc.core_properties.title = "金箔非遗项目四岗位7周任务与验收表 V0.1"
    doc.core_properties.subject = "广东省攀登计划申报前7周学生执行与验收基线"
    doc.core_properties.author = "项目团队"
    doc.core_properties.keywords = "金箔, 非遗, 四岗位, 七周排期, 任务, 验收"


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 12.5, bold=True, color=INK if level == 1 else GOLD_DARK)
    return paragraph


def add_table(doc: Document, headers, rows, weights, font_size=8.0, zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        fill_cell(table.rows[0].cells[idx], header, bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[idx], GOLD_DARK)
        table.rows[0].cells[idx].paragraphs[0].paragraph_format.keep_with_next = True
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            align = WD_ALIGN_PARAGRAPH.CENTER if cidx in (0, 5, 6) else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(cell, str(value), bold=(cidx == 0), color=GOLD_DARK if cidx == 0 else INK, size=font_size, align=align)
            if zebra and ridx % 2 == 1:
                set_cell_shading(cell, CREAM_LIGHT)

    section = doc.sections[-1]
    content_width = section_content_width_dxa(section)
    table_width = content_width - 120
    widths = column_widths_from_weights(weights, table_width)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=table_width,
        indent_dxa=120,
        cell_margins_dxa={"top": 90, "bottom": 90, "start": 120, "end": 120},
    )
    set_table_borders(table, color=GRID, size="5")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_opening(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(6)
    kicker.paragraph_format.space_after = Pt(2)
    set_run_font(kicker.add_run("广东省攀登计划 · 学生执行文件"), size=9.5, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    set_run_font(title.add_run("金箔非遗项目四岗位7周任务与验收表"), size=24, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    set_run_font(subtitle.add_run("V0.1｜从岗位选择到P0申报演示验收"), size=11.5, color=MUTED)

    add_table(
        doc,
        ["执行项", "当前定义"],
        [
            ("执行周期", "2026年7月17日—9月3日，共7周；官方节点变化时整体平移"),
            ("团队结构", "4名学生、4个主责岗位；每个岗位设置可接手的备份协作人"),
            ("核心目标", "完成P0三步体验、真实内容、埋点、真机测试和申报演示证据"),
            ("验收方式", "每周提交可检查交付物，由协作岗位复核，指导教师完成里程碑验收"),
        ],
        [1.4, 8.6],
        font_size=8.4,
        zebra=False,
    )

    add_callout(
        doc,
        "执行规则｜",
        "状态统一使用“未开始、进行中、待验收、已通过、需返工、已阻塞”。工艺真实性、素材授权和知情同意属于否决性条件，不以进度压力降低要求。",
        fill=CREAM,
    )

    add_heading(doc, "岗位边界与固定协作", 2)
    add_table(
        doc,
        ["岗位", "主责", "固定协作与备份"],
        [
            ("1 产品与前端", "原型、页面、视频呈现、结果页、内容入口和前端埋点", "与岗位2共同维护主流程；与岗位3完成内容上架"),
            ("2 3D与技术", "XR-Frame、3D交互、CloudBase、性能、权限、部署和数据", "与岗位1共同开发；与岗位4完成埋点核验"),
            ("3 内容与档案", "访谈、拍摄、授权、内容台账、专业审核和传播素材", "与岗位1上架内容；与岗位4支持活动和访谈"),
            ("4 调研与实践", "研究、同意、问卷、测试、活动、数据分析和证据归档", "与岗位2复算数据；与岗位3开展校园实践"),
        ],
        [1.3, 4.4, 4.3],
        font_size=8.1,
    )


def add_week_page(doc: Document, week: dict) -> None:
    doc.add_page_break()
    add_heading(doc, f"{week['week']}｜{week['dates']}", 1)
    add_callout(doc, "本周验收门｜", week["goal"], fill=CREAM_LIGHT)
    add_table(
        doc,
        ["岗位", "主责任务", "协作／前置依赖", "本周交付物", "验收标准", "验收人", "状态"],
        week["rows"],
        [1.1, 2.9, 2.45, 2.4, 3.0, 1.25, 1.0],
        font_size=7.8,
    )
    add_body(
        doc,
        "周验收记录：主责人提交交付物与证据；协作人确认接口和交接；指导教师记录“通过、返工或阻塞”结论及下一步。",
        after=0,
    )


def add_final_acceptance(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "P0最终验收与交接清单", 1)
    add_callout(
        doc,
        "冻结条件｜",
        "只有下列验收域全部通过，才可将版本标记为P0申报演示候选版。目标值、测试结果和未来计划必须在材料中分别表述。",
        fill=CREAM,
    )
    add_table(
        doc,
        ["验收域", "最终检查", "必须提交的证据", "责任岗位", "结论"],
        [
            ("产品流程", "首页至结果页无需注册可完成；HOME／EXP／INT／RES逐项通过", "体验版二维码、截图、验收记录", "岗位1／2", "待验收"),
            ("3D与降级", "三步完整；至少1步高保真；加载、弱网和失败降级有效", "6台设备记录、性能与降级报告", "岗位2", "待验收"),
            ("内容真实性", "视频、术语、人物和工艺经审核，简化模拟有说明", "内容台账、审核与版本记录", "岗位3", "待验收"),
            ("授权合规", "公开素材均在授权范围内，未授权素材不进入公开版本", "授权索引、素材ID和许可文件", "岗位3", "待验收"),
            ("数据埋点", "事件不漏报、不重复；测试数据可识别；指标可复算", "埋点字典、事件核验和漏斗表", "岗位2／4", "待验收"),
            ("研究与隐私", "同意文本、研究编号、三表分离、退出和限制访问可用", "问卷、同意文本、数据字典", "岗位4", "待验收"),
            ("用户测试", "完成30—50人测试并形成问题、修复和数据摘要", "招募、问卷、访谈和闭环记录", "岗位4", "待验收"),
            ("部署与备份", "云端、本地和离线演示方案可用；版本可恢复", "部署记录、备份包、恢复说明", "岗位2", "待验收"),
            ("申报证据", "产品、内容、数据、合作与实践材料可追溯", "证据目录、PPT页和演示脚本", "全体", "待验收"),
        ],
        [1.3, 3.2, 3.0, 1.25, 1.0],
        font_size=7.9,
    )
    add_heading(doc, "交接要求", 2)
    add_body(doc, "每个岗位必须提交最新版文件、源文件位置、账号与权限说明、已知问题、恢复方式和备份人确认。仅提交截图或口头说明不视为完成交接。")
    add_callout(
        doc,
        "变更规则｜",
        "新增任务必须说明对应的PRD需求、影响周次、主责人和验收标准；未通过范围评审的P1/P2功能不得进入P0工期。",
        fill=SOFT_GRAY,
    )


def validate(doc: Document) -> None:
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    for week in range(1, 8):
        if f"第{week}周" not in text:
            raise AssertionError(f"Missing week {week}")
    role_week_rows = 0
    for table in doc.tables:
        if [cell.text for cell in table.rows[0].cells][:2] == ["岗位", "主责任务"]:
            role_week_rows += len(table.rows) - 1
            for row in table.rows[1:]:
                if not row.cells[3].text.strip() or not row.cells[4].text.strip():
                    raise AssertionError("A role-week row lacks a deliverable or acceptance standard")
    if role_week_rows != 28:
        raise AssertionError(f"Expected 28 role-week rows, found {role_week_rows}")
    forbidden = ["TODO", "TBD", "待填写", "{{", "}}"]
    hits = [item for item in forbidden if item in text]
    if hits:
        raise AssertionError(f"Placeholder markers found: {hits}")


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_opening(doc)
    for week in WEEKS:
        add_week_page(doc, week)
    add_final_acceptance(doc)
    validate(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
