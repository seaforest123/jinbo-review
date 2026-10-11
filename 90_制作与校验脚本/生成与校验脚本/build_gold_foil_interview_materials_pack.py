from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_gold_foil_project_brief_docx import (  # noqa: E402
    CREAM,
    CREAM_LIGHT,
    GOLD,
    GOLD_DARK,
    GRID,
    INK,
    MUTED,
    WHITE,
    add_field,
    fill_cell,
    prevent_row_split,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_style_font,
    set_table_borders,
)

TABLE_GEOMETRY_DIR = Path(
    r"C:\Users\Andrew\.codex\plugins\cache\openai-primary-runtime\documents\26.715.12143\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_GEOMETRY_DIR))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


OUT_DIR = ROOT / "docs" / "gold-foil-heritage" / "首次专业访谈材料包_V0.1_20260719"
VERSION = "V0.1"
DATE = "2026年7月19日"
PROJECT = "金箔非遗数字体验与活化传播项目"
PALE_GOLD = "F7F0DF"
SOFT_GOLD = "E7D7AB"
SOFT_RED = "F7E9E6"
SOFT_GREEN = "EAF2E8"

FILES = [
    "00_首次专业访谈材料包总目录与使用顺序_V0.1_20260719.docx",
    "01_金箔非遗数字体验项目一页纸_V0.1_20260719.docx",
    "02_三分钟轻量工艺体验纸面原型_V0.1_20260719.docx",
    "03_纸面原型演示脚本与专家反馈表_V0.1_20260719.docx",
    "04_首次专业访谈提纲与60分钟主持脚本_V0.1_20260719.docx",
    "05_拍摄录音及素材使用授权确认单_V0.1_20260719.docx",
    "06_访谈对象身份与联系人信息确认表_V0.1_20260719.docx",
    "07_时间地点确认及微信约访文本_V0.1_20260719.docx",
    "08_现场设备与记录工具检查清单_V0.1_20260719.docx",
    "09_访谈纪要关键结论与会后确认单_V0.1_20260719.docx",
]


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_row_min_height(row, mm: float) -> None:
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def configure_document(doc: Document, code: str, short_title: str) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.72)
    section.footer_distance = Cm(0.68)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.0, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.18

    for level, size in ((1, 16.5), (2, 12.6), (3, 11.0)):
        style = doc.styles[f"Heading {level}"]
        set_style_font(style, size, INK if level == 1 else GOLD_DARK, bold=True)
        style.paragraph_format.space_before = Pt(8 if level == 1 else 6)
        style.paragraph_format.space_after = Pt(5 if level == 1 else 3)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    clear_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run(PROJECT), size=8.0, bold=True, color=MUTED)
    set_run_font(header.add_run(f"    |    {code} {short_title}"), size=8.0, color=MUTED)

    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run(f"{VERSION} · {DATE}    第 "), size=8.0, color=MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.add_run(" 页"), size=8.0, color=MUTED)

    doc.core_properties.title = f"{short_title} {VERSION}"
    doc.core_properties.subject = "金箔非遗项目首次专业访谈材料包"
    doc.core_properties.author = "金箔非遗项目团队"
    doc.core_properties.keywords = "金箔, 非遗, 专业访谈, 攀登计划, 纸面原型, 授权"


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    if getattr(doc, "_gold_page_break_pending", False):
        p.paragraph_format.page_break_before = True
        setattr(doc, "_gold_page_break_pending", False)
    set_run_font(p.add_run(text), size={1: 16.5, 2: 12.6, 3: 11.0}[level], bold=True,
                 color=INK if level == 1 else GOLD_DARK)
    return p


def add_body(doc: Document, text: str, *, size=10.0, bold=False, color=INK, after=4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_callout(doc: Document, label: str, text: str, *, fill=CREAM_LIGHT):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(7)
    p.paragraph_format.right_indent = Pt(5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.16
    set_paragraph_shading(p, fill)
    set_run_font(p.add_run(label), size=9.8, bold=True, color=GOLD_DARK)
    set_run_font(p.add_run(text), size=9.8, color=INK)
    return p


def add_title(doc: Document, code: str, title: str, purpose: str) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(5)
    kicker.paragraph_format.space_after = Pt(2)
    set_run_font(kicker.add_run(f"首次专业访谈材料包 · {code}"), size=9.2, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(title), size=21, bold=True, color=INK)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(7)
    set_run_font(meta.add_run(f"{VERSION}｜{DATE}｜{PROJECT}"), size=9.0, color=MUTED)
    add_callout(doc, "用途｜", purpose, fill=CREAM)


def add_page_break(doc: Document) -> None:
    # Applying page_break_before to the next heading avoids a blank page when
    # the preceding table already reaches the bottom of the current page.
    remove_trailing_blank_paragraphs(doc)
    setattr(doc, "_gold_page_break_pending", True)


def add_table(doc: Document, headers, rows, weights, *, font_size=8.6, center_cols=(), zebra=True,
              min_row_mm=None, header_fill=GOLD_DARK):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        fill_cell(table.rows[0].cells[idx], str(header), bold=True, color=WHITE, size=8.6,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[idx], header_fill)
        table.rows[0].cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if min_row_mm is not None:
            set_row_min_height(row, min_row_mm)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            align = WD_ALIGN_PARAGRAPH.CENTER if cidx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(cell, str(value), size=font_size, align=align, color=INK)
            if zebra and ridx % 2:
                set_cell_shading(cell, CREAM_LIGHT)
    section = doc.sections[-1]
    content_width = section_content_width_dxa(section)
    table_width = content_width - 80
    widths = column_widths_from_weights(weights, table_width)
    apply_table_geometry(table, widths, table_width_dxa=table_width, indent_dxa=105,
                         cell_margins_dxa={"top": 80, "bottom": 80, "start": 105, "end": 105})
    set_table_borders(table, color=GRID, size="5")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return table


def add_form(doc: Document, rows, *, weights=(1.7, 5.3), min_row_mm=9, font_size=9.0):
    return add_table(doc, ["字段", "填写／确认内容"], rows, weights, font_size=font_size,
                     center_cols=(0,), zebra=False, min_row_mm=min_row_mm)


def add_checklist(doc: Document, rows, *, detail_header="检查内容", owner=True, min_row_mm=8):
    if owner:
        normalized = [tuple(row) + ("",) if len(row) == 4 else tuple(row) for row in rows]
        return add_table(doc, ["完成状态", detail_header, "负责人", "异常说明", "补救措施"], normalized,
                         [0.55, 3.45, 0.95, 1.35, 1.4], font_size=8.0, center_cols=(0, 2),
                         zebra=False, min_row_mm=min_row_mm)
    return add_table(doc, ["状态", detail_header, "说明／记录"], rows,
                     [0.7, 4.2, 2.6], font_size=8.6, center_cols=(0,), zebra=False,
                     min_row_mm=min_row_mm)


def remove_trailing_blank_paragraphs(doc: Document) -> None:
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        paragraph = doc.paragraphs[-1]
        parent = paragraph._p.getparent()
        if parent is None:
            break
        parent.remove(paragraph._p)


def save(doc: Document, filename: str) -> None:
    remove_trailing_blank_paragraphs(doc)
    doc.save(OUT_DIR / filename)


def base_doc(code: str, title: str) -> Document:
    doc = Document()
    configure_document(doc, code, title)
    return doc


def build_00() -> None:
    title = "首次专业访谈材料包总目录与使用顺序"
    doc = base_doc("00", title)
    doc.sections[0].bottom_margin = Cm(1.1)
    add_title(doc, "00", title, "整套材料的入口、使用顺序和归档索引；由指导教师或项目统筹在访谈前建立并在访谈后核验。")
    add_callout(doc, "材料构成｜", "9份独立使用文件＋1份总目录；每份均可独立打开、打印和填写。", fill=PALE_GOLD)
    add_heading(doc, "一、使用原则", 2)
    add_checklist(doc, [
        ("□", "未经专业确认的工艺动作、术语、顺序、工具材料，一律标注“候选方案／待专业确认”。", "已知信息也注明来源"),
        ("□", "录音、录像、拍照、姓名、肖像、作品和公开传播均逐项确认。", "不同用途不得互相替代"),
        ("□", "历史照片、作品图片和第三方材料只登记线索，不因本次访谈自动获得授权。", "权属清楚后另行办理"),
        ("□", "只有已确认且在授权范围内的内容才能进入公开小程序、申报附件、答辩或短视频。", "保留版本与审核记录"),
    ], owner=False, min_row_mm=10)
    add_heading(doc, "二、文件总目录", 2)
    rows = [
        ("00", "总目录与使用顺序", "教师／统筹", "前—中—后", "完整性确认与总归档"),
        ("01", "项目一页纸", "教师／岗位3", "会前发送、开场", "发送记录／现场签收"),
        ("02", "三分钟纸面原型", "岗位1", "会前准备、现场评审", "逐屏批注原件"),
        ("03", "演示脚本与反馈表", "岗位2＋岗位1", "现场演示", "专家反馈与优先级"),
        ("04", "60分钟主持脚本", "岗位3", "现场主持", "问题记录与时间标记"),
        ("05", "拍摄录音及素材授权", "岗位3＋教师", "记录前、结束前", "签署件／复核要求"),
        ("06", "身份与联系人确认", "岗位3", "会前或现场", "核验来源与联系偏好"),
        ("07", "时间地点及微信文本", "岗位3", "约访与提醒", "沟通截图／确认信息"),
        ("08", "设备与记录检查清单", "岗位2＋全体", "会前、进场、离场", "完成状态与补救"),
        ("09", "纪要与会后确认单", "岗位4＋全体", "会后24小时内", "确认稿及回复证据"),
    ]
    rows = [tuple(row) + ("现场填写／按表内勾选", "□") for row in rows]
    add_table(doc, ["编号", "材料", "主责", "使用时点", "归档位置／结果", "填写要求", "完成状态"], rows,
              [0.45, 1.65, 0.85, 0.95, 1.55, 1.35, 0.7], font_size=7.3, center_cols=(0, 2, 3, 6), min_row_mm=8)

    add_page_break(doc)
    add_heading(doc, "三、访谈前—访谈中—访谈后使用顺序", 1)
    add_table(doc, ["阶段", "顺序", "动作", "使用材料", "完成证据"], [
        ("访谈前", "1", "发送项目一页纸，完成正式约访和时间地点确认。", "01、07", "对方确认信息"),
        ("访谈前", "2", "填写已知身份信息；不得把未核实资格写成事实。", "06", "信息来源记录"),
        ("访谈前", "3", "准备双路线纸面原型、主持脚本、授权文件和备份设备。", "02—05、08", "打印件与检查结果"),
        ("访谈中", "4", "说明边界；取得相应同意后再录音、录像或拍照。", "01、05", "勾选／签名"),
        ("访谈中", "5", "演示原型，优先完成五个必答问题；负面意见记录原话。", "02—04", "逐屏反馈／原话"),
        ("访谈中", "6", "结束前确认路线、真实性红线、授权、审核人和下次沟通。", "04—06", "结论与联系人"),
        ("访谈后24小时内", "7", "形成纪要，区分已确认、待确认、被否定和需补证据。", "09", "纪要初稿"),
        ("访谈后", "8", "向专业老师发送简短确认稿；仅使用获确认且获授权内容。", "09", "回复截图／邮件"),
        ("访谈后", "9", "把确认结论转为四岗位第2周行动并分类归档。", "00、09", "任务表与资料路径"),
    ], [1.0, 0.55, 3.1, 1.0, 1.85], font_size=8.2, center_cols=(0, 1, 3), min_row_mm=9)
    add_heading(doc, "四、必带材料与可选材料", 2)
    add_table(doc, ["类别", "清单", "数量／位置", "确认"], [
        ("必带纸质", "01、02、03、04、05、06、08；签字笔；便签；夹板", "每类建议2份", "□"),
        ("必带电子", "十份Word、纸面原型PDF、离线备份、统一文件名", "手机＋电脑／U盘", "□"),
        ("必带设备", "主录音设备、备用手机、充电宝、充电线", "访谈前实测", "□"),
        ("可选", "已有项目简报、PRD、团队分工表、参考图片线索", "不替代本材料包", "□"),
    ], [1.0, 4.1, 1.7, 0.7], font_size=8.6, center_cols=(0, 2, 3), zebra=False, min_row_mm=10)

    add_page_break(doc)
    add_heading(doc, "五、资料完整性与归档确认", 1)
    add_form(doc, [
        ("项目联系人", "姓名：________________  手机／微信：________________"),
        ("访谈对象", "姓名／称呼：________________  机构：________________"),
        ("访谈时间地点", "日期：____年__月__日  时间：____—____  地点：________________"),
        ("总归档位置", "____________________________________________________________"),
        ("资料命名规则", "日期_对象_材料类型_版本；原始素材与公开版本分开保存"),
    ], min_row_mm=10)
    add_checklist(doc, [("□", f"{name}", "", "") for name in FILES], detail_header="文件名／归档项", min_row_mm=7)
    add_callout(doc, "完整性结论｜", "□ 完整，可进入会后确认　□ 有缺项，已记录补办责任人和期限　□ 暂停公开使用", fill=SOFT_GREEN)
    add_form(doc, [("核验人", "________________"), ("核验日期", "____年__月__日"), ("补办说明", "____________________________________________________________")], min_row_mm=8)
    save(doc, FILES[0])


def build_01() -> None:
    title = "金箔非遗数字体验项目一页纸"
    doc = base_doc("01", title)
    section = doc.sections[0]
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.15)
    doc.styles["Normal"].paragraph_format.space_after = Pt(2)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    h2 = doc.styles["Heading 2"]
    h2.paragraph_format.space_before = Pt(3)
    h2.paragraph_format.space_after = Pt(1.5)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(1)
    kicker.paragraph_format.space_after = Pt(1)
    set_run_font(kicker.add_run("首次专业访谈材料包 · 01"), size=8.6, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_run_font(p.add_run(title), size=18.5, bold=True, color=INK)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(3)
    set_run_font(meta.add_run(f"{VERSION}｜{DATE}｜{PROJECT}"), size=8.2, color=MUTED)
    purpose = doc.add_paragraph()
    purpose.paragraph_format.space_after = Pt(3)
    set_paragraph_shading(purpose, CREAM)
    set_run_font(purpose.add_run("用途｜"), size=8.7, bold=True, color=GOLD_DARK)
    set_run_font(purpose.add_run("供专业老师会前或开场快速理解项目；不把候选工艺路线写成既定结论。"), size=8.7, color=INK)
    add_heading(doc, "我们为什么做", 2)
    add_body(doc, "团队希望回应金箔相关非遗“公众接触少、青年理解门槛高、数字内容零散”的现实问题，以轻量互动建立第一次接触，再把愿意深入的用户引向人物故事、专业知识与线下体验。现状判断与具体表述均待专业访谈核实。", size=8.9, after=2)
    add_heading(doc, "首阶段用户与场景", 2)
    add_table(doc, ["核心用户", "典型场景", "首轮目标"], [
        ("广东高校学生，重点为18—24岁青年", "校园传播、课堂／社团体验、展览扫码、短内容引流", "愿意进入、3分钟内完成、愿意继续了解或参与线下活动"),
    ], [2.0, 2.6, 2.9], font_size=8.2, center_cols=(0,), zebra=False, min_row_mm=10)
    add_heading(doc, "产品定位", 2)
    add_body(doc, "主线｜3分钟轻量工艺体验＋可选深度学习。互动模拟是核心入口，文化内容提供理解，3D／AR是候选增强手段，不先于内容真实性和基础体验。", size=8.7, bold=True, color=GOLD_DARK, after=2)
    add_table(doc, ["层级", "首版作用", "边界"], [
        ("互动模拟", "用三个经确认的关键节点形成轻量操作与即时反馈", "不复制完整工艺，不暗示掌握真实技艺"),
        ("文化内容", "补充人物、工艺背景、价值与线下连接", "未经审核不公开；来源可追溯"),
        ("3D／AR候选", "增强材料、工具或结果的观察与空间呈现", "有必要且成本可控时采用，不作为首版成败前提"),
    ], [1.25, 3.25, 3.0], font_size=7.9, center_cols=(0,), min_row_mm=7.5)
    add_heading(doc, "本次首次访谈希望解决的五个问题", 2)
    add_table(doc, ["1", "2", "3", "4", "5"], [[
        "公开身份应如何准确表述？", "首版更适合制箔、贴箔还是第三种路线？", "三分钟体验应保留哪三个关键节点？", "哪些表述、动作或画面会失真／不宜公开？", "哪些记录与使用范围可以授权，哪些需发布前复核？"
    ]], [1.5]*5, font_size=7.25, center_cols=(0,1,2,3,4), zebra=False, min_row_mm=14)
    add_body(doc, "依次覆盖：身份称谓、首版路线、三个互动节点、真实性红线、拍摄与素材使用范围。", size=8.0, color=MUTED, after=2)
    add_heading(doc, "项目边界与团队协作", 2)
    add_table(doc, ["明确不做", "四岗位并行", "期待合作"], [
        ("不教授完整真实工艺；不替代线下传承；未经审核不公开；第三方材料不默认获权。", "岗位1产品与交互｜岗位2技术开发｜岗位3专业对接与传播｜岗位4研究与证据。", "请专业老师提供路线判断、真实性红线、公开边界、审核机制及后续资料线索。"),
    ], [2.55, 2.45, 2.5], font_size=7.65, min_row_mm=11)
    save(doc, FILES[1])


def add_screen(doc: Document, number: int, name: str, goal: str, content_rows, caution: str) -> None:
    add_heading(doc, f"界面{number}｜{name}", 1)
    add_callout(doc, "界面目标｜", goal, fill=PALE_GOLD)
    add_table(doc, ["区域", "纸面展示内容", "专业确认"], content_rows,
              [1.2, 4.3, 2.0], font_size=8.8, center_cols=(0,), zebra=False, min_row_mm=12)
    add_callout(doc, "边界提示｜", caution, fill=SOFT_RED)
    add_table(doc, ["评审选择", "具体意见／需补证据"], [
        ("□ 保留　□ 修改　□ 删除　□ 待补证据", ""),
    ], [2.8, 4.7], font_size=9.0, center_cols=(0,), zebra=False, min_row_mm=20)


def build_02() -> None:
    title = "三分钟轻量工艺体验纸面原型"
    doc = base_doc("02", title)
    add_title(doc, "02", title, "替代尚未完成的软件Demo，让专业老师对五个界面、候选路线和信息边界进行具体评审。")
    add_callout(doc, "重要说明｜", "本原型只展示信息结构和交互逻辑。所有工艺动作、术语、顺序、工具材料均为“候选方案／待专业确认”，不得据此制作公开内容。", fill=SOFT_RED)
    add_form(doc, [("评审人", "________________"), ("评审日期", "____年__月__日"), ("记录人", "________________")], min_row_mm=9)
    add_heading(doc, "总流程草图", 2)
    add_table(doc, ["进入体验", "候选路线", "三步互动", "体验结果", "可选深入"], [[
        "知道这是一段数字轻体验", "选择／否定／提出第三种", "操作3个待确认节点", "看到数字结果与文化提示", "人物／知识／活动／资料"
    ]], [1.5]*5, font_size=8.0, center_cols=(0,1,2,3,4), zebra=False, min_row_mm=17)

    add_page_break(doc)
    add_screen(doc, 1, "进入体验", "让用户愿意开始，并清楚这不是完整技艺教学。", [
        ("标题", "金箔数字轻体验（暂名，可修改）", "□ 保留 □ 修改"),
        ("一句话", "在约3分钟内认识一个经专业确认的关键工艺片段。", "是否准确：________"),
        ("主按钮", "开始体验", "是否有吸引力：________"),
        ("提示", "数字体验不替代真实工艺学习与线下传承。", "是否充分：________"),
    ], "不得以“学会金箔工艺”“成为传承人”等夸大措辞吸引点击。")

    add_page_break(doc)
    add_screen(doc, 2, "候选路线", "在专业判断前并列展示两条候选路线，并允许提出第三种方案。", [
        ("候选A", "制箔候选路线｜适合首版吗？为什么？", "□ 适合 □ 不适合 □ 待研究"),
        ("候选B", "贴箔候选路线｜适合首版吗？为什么？", "□ 适合 □ 不适合 □ 待研究"),
        ("候选C", "第三种路线：____________________________", "提出人：____________"),
        ("判断维度", "真实性、代表性、3分钟可表达性、互动性、素材可获得性、公开风险", "最重要维度：________"),
    ], "本页不代表团队已选定制箔或贴箔；如暂不能确定，两条路线均保留为候选。")

    add_page_break(doc)
    add_screen(doc, 3, "三步互动", "确认首版应保留的三个关键互动节点及正确顺序。", [
        ("节点1", "名称／动作：________________　工具材料：________________", "顺序：__ 风险：________"),
        ("节点2", "名称／动作：________________　工具材料：________________", "顺序：__ 风险：________"),
        ("节点3", "名称／动作：________________　工具材料：________________", "顺序：__ 风险：________"),
        ("交互反馈", "用户做对、做错或停顿时，系统应如何提示：________________", "表述审核人：________"),
    ], "节点必须由专业老师确认；数字化可简化操作，但不得颠倒因果、误用术语或把危险／保密环节游戏化。")

    add_page_break(doc)
    add_screen(doc, 4, "体验结果", "给予完成感，同时把数字结果与真实工艺能力明确区分。", [
        ("数字结果", "数字作品／过程回放／文化卡片（候选，可多选）", "□ 保留 □ 修改"),
        ("必备提示", "你完成的是数字化轻量体验，不代表已经掌握真实工艺。", "□ 准确 □ 修改"),
        ("文化信息", "一条经审核的工艺价值或人物信息：________________", "来源：____________"),
        ("下一步", "继续了解／预约线下活动／分享（按授权和功能范围选择）", "优先：____________"),
    ], "不得用等级、证书或称号让用户误认为获得专业能力、传承资格或权威认证。")

    add_page_break(doc)
    add_screen(doc, 5, "可选深入了解", "让有兴趣的用户自主进入深度内容，不阻塞三分钟主线。", [
        ("人物故事", "公开身份、经历与观点；内容待本人审核。", "□ 可做 □ 暂缓"),
        ("工艺知识", "经确认的术语、背景、工具材料与延伸阅读。", "□ 可做 □ 暂缓"),
        ("线下连接", "讲座、工作坊、展览或预约信息。", "□ 可做 □ 暂缓"),
        ("资料页", "来源、审核说明、版本与更正渠道。", "□ 必须 □ 可选"),
    ], "深度内容不是知识测评入口；手机号等个人信息只在明确场景、用途和同意后收集。")
    add_heading(doc, "整体评审结论", 2)
    add_form(doc, [
        ("建议首版路线", "□ 制箔候选　□ 贴箔候选　□ 第三种：____________　□ 暂不确定"),
        ("最需修改界面", "第__屏；原因：____________________________________________"),
        ("真实性红线", "____________________________________________________________"),
        ("下一次复核材料", "____________________________________________________________"),
    ], min_row_mm=11)
    save(doc, FILES[2])


def build_03() -> None:
    title = "纸面原型演示脚本与专家反馈表"
    doc = base_doc("03", title)
    add_title(doc, "03", title, "统一5—10分钟演示方式，记录专家第一反应、逐屏判断、路线理由和改进优先级。")
    add_form(doc, [("演示人", "________________"), ("记录人", "________________"),
                   ("专业老师", "________________"), ("日期", "____年__月__日"),
                   ("纸面原型版本", "________________")], min_row_mm=7)
    add_heading(doc, "一、演示前30秒", 2)
    add_callout(doc, "主持话术｜", "老师您好，目前尚无可运行Demo，本次展示的是纸面原型，不是已经定稿的软件方案。制箔、贴箔两条路线均为候选方案，具体工艺内容待专业确认。我们希望您直接指出不准确、不合适或不应公开之处，负面反馈只记录原话，不在现场辩解或说服。", fill=PALE_GOLD)
    add_form(doc, [("专家第一反应", "____________________________________________________________\n____________________________________________________________"),
                   ("记录方式", "□ 原话　□ 概括（需会后确认）　时间点：________")], min_row_mm=14)
    add_heading(doc, "二、5—10分钟统一演示脚本", 2)
    add_table(doc, ["用时", "界面", "演示话术与动作", "不得预设"], [
        ("约1分钟", "1 进入体验", "说明3分钟定位、目标用户和“不替代真实工艺”。询问入口是否愿意点开。", "不承诺学会技艺"),
        ("约2分钟", "2 候选路线", "并列展示制箔、贴箔和第三种选择，请专家先判断路线，再说明理由。", "不暗示偏好"),
        ("约2分钟", "3 三步互动", "展示三个空节点，请专家填写名称、顺序、工具材料和不能简化之处。", "不自填动作"),
        ("约1分钟", "4 体验结果", "展示数字结果和能力边界提示，请判断是否准确、是否会误导。", "不设置权威认证"),
        ("约1分钟", "5 可选深入", "展示人物、知识、线下活动和资料页，请判断优先级与公开边界。", "不强迫学习测评"),
        ("1—3分钟", "收束", "复述路线、三个节点、真实性红线、优先修改项，请专家纠正。", "不当场争辩"),
    ], [0.9, 1.15, 4.2, 1.25], font_size=8.1, center_cols=(0,1), min_row_mm=9)

    add_page_break(doc)
    add_heading(doc, "三、逐屏反馈记录", 1)
    rows = []
    for i, name in enumerate(["进入体验", "候选路线", "三步互动", "体验结果", "可选深入"], 1):
        rows.append((f"{i} {name}", "□ 保留 □ 修改 □ 删除 □ 待补证据", "", ""))
    add_table(doc, ["界面", "判断", "专家意见／原话", "证据或复核人"], rows,
              [1.2, 2.0, 3.2, 1.3], font_size=8.5, center_cols=(0,1), zebra=False, min_row_mm=21)
    add_callout(doc, "记录纪律｜", "尽量记录原话；无法完整记录时标注“概括”，会后写入确认稿。专家否定某项时，不现场说服，追问适用条件和可替代方案。", fill=SOFT_RED)

    add_page_break(doc)
    add_heading(doc, "四、工艺路线与三个节点", 1)
    add_form(doc, [
        ("建议首版路线", "□ 制箔候选　□ 贴箔候选　□ 第三种：____________　□ 暂不确定"),
        ("专业理由", "____________________________________________________________\n____________________________________________________________"),
        ("不选另一条路线的原因", "____________________________________________________________"),
        ("判断适用范围", "□ 本项目首版　□ 面向高校学生　□ 仅当前材料条件　□ 其他：________"),
    ], min_row_mm=14)
    add_table(doc, ["节点", "正确名称／顺序", "可数字化的部分", "不可简化／风险"], [
        ("1", "", "", ""), ("2", "", "", ""), ("3", "", "", "")
    ], [0.65, 2.2, 2.3, 2.55], font_size=8.7, center_cols=(0,), zebra=False, min_row_mm=20)
    add_heading(doc, "五、真实性风险与改进优先级", 2)
    add_table(doc, ["优先级", "问题／风险", "建议修改", "确认方式"], [
        ("P0 必须先改", "", "", ""), ("P1 首版应改", "", "", ""), ("P2 后续优化", "", "", "")
    ], [1.1, 2.5, 2.6, 1.5], font_size=8.6, center_cols=(0,), zebra=False, min_row_mm=17)
    add_heading(doc, "六、会后再次确认", 2)
    add_checklist(doc, [
        ("□", "身份与公开称谓", "", ""), ("□", "路线与三个节点", "", ""),
        ("□", "真实性红线／不可公开内容", "", ""), ("□", "拟公开文字、图片或视频版本", "", ""),
        ("□", "第三方资料权属线索", "", ""),
    ], detail_header="复核事项", min_row_mm=8)
    save(doc, FILES[3])


def build_04() -> None:
    title = "首次专业访谈提纲与60分钟主持脚本"
    doc = base_doc("04", title)
    add_title(doc, "04", title, "支持主持人控制60分钟流程；时间不足时优先完成五个必答问题和授权边界。")
    add_form(doc, [("访谈对象", "________________"), ("主持人", "________________"),
                   ("记录人／设备", "________________／________________"),
                   ("日期地点", "____年__月__日　________________"),
                   ("纸面原型版本", "________________")], min_row_mm=8)
    add_heading(doc, "一、60分钟总流程", 2)
    add_table(doc, ["时间", "环节", "主持动作", "最低完成结果"], [
        ("0—5分钟", "介绍与授权", "说明项目、访谈用途、记录方式；按05逐项取得同意。", "知道可拒绝；记录范围明确"),
        ("5—15分钟", "原型与第一反应", "按03展示五屏纸面原型，不预设路线。", "第一反应＋明显错误"),
        ("15—40分钟", "第一层必答", "完成身份、路线、三个节点、真实性红线、拍摄与使用范围。", "五项均有结论或待确认"),
        ("40—55分钟", "第二层选答", "追问交互细节、结果提示、审核机制、项目期待。", "记录优先级与联系人"),
        ("55—60分钟", "复述与收尾", "复述结论；确认授权、审核、下次沟通及会后确认方式。", "对方纠正／认可复述"),
    ], [1.05, 1.25, 3.75, 2.0], font_size=8.25, center_cols=(0,1), min_row_mm=10)
    add_callout(doc, "时间不足降级｜", "立即停止第二层追问，优先完成以下五个问题和授权确认；身份资料权属、细节解释和下次安排改为会后书面补充。", fill=SOFT_RED)

    add_page_break(doc)
    add_heading(doc, "二、第一层：五个必答问题（15—40分钟）", 1)
    questions = [
        ("1 身份称谓", "我们在内部记录和公开材料中，应怎样准确介绍您的姓名、机构、职务以及与金箔工艺的关系？哪些称谓不能使用？"),
        ("2 首版路线", "面向广东高校学生的3分钟体验，制箔、贴箔或其他路线中，哪一条更适合首版？判断理由和适用条件是什么？"),
        ("3 三个节点", "如果只保留三个互动节点，应分别是什么、正确顺序怎样、哪些工具材料或动作不能被简化？"),
        ("4 真实性红线", "哪些术语、画面、动作、故事或结果提示容易造成误解？哪些内容不宜记录或公开？"),
        ("5 记录与使用", "本次录音、录像、拍照以及姓名、肖像、声音、工艺和作品，分别允许用于哪些场景？哪些需发布前复核？"),
    ]
    for label, q in questions:
        add_heading(doc, label, 2)
        add_body(doc, q, bold=True, size=9.6)
        add_table(doc, ["原话／结论", "追问与条件", "状态"], [("", "", "□ 已确认\n□ 待确认")],
                  [3.8, 2.5, 1.2], font_size=8.7, center_cols=(2,), zebra=False, min_row_mm=24)

    add_page_break(doc)
    add_heading(doc, "三、第二层：现场选答问题（40—55分钟）", 1)
    add_table(doc, ["主题", "可选问题", "记录重点"], [
        ("交互细节", "用户可以做什么？哪些反馈既有趣又不失真？做错时如何提示？", "动作、反馈、禁区、优先级"),
        ("结果页", "“数字体验不代表掌握真实工艺”的提示是否准确？怎样更自然？", "建议原话、不得使用的称谓"),
        ("深度内容", "人物故事、工艺知识、线下活动、资料页应优先做哪项？", "来源、审核、可公开范围"),
        ("3D／AR", "哪些对象值得3D／AR呈现？是否真的有助理解？有哪些成本或失真风险？", "必要性、素材、设备、替代方案"),
        ("审核机制", "由谁审核文字、图片、视频和交互？通过什么渠道？一般需要多久？", "审核人、周期、版本留痕"),
        ("项目期待", "您希望项目对传承、传播或线下活动产生什么实际帮助？怎样判断有效？", "成果形式、可观察指标"),
    ], [1.1, 4.25, 2.15], font_size=8.35, center_cols=(0,), min_row_mm=12)
    add_heading(doc, "现场选答记录", 2)
    add_table(doc, ["主题", "专家意见／原话", "下一步"], [("", "", ""), ("", "", ""), ("", "", "")],
              [1.15, 4.8, 1.55], font_size=8.7, zebra=False, min_row_mm=17)

    add_page_break(doc)
    add_heading(doc, "四、第三层：会后书面确认", 1)
    add_checklist(doc, [
        ("□", "历史照片、作品图片、旧视频等资料由谁持有？谁有权授权？", "", ""),
        ("□", "公开身份和资格表述是否有可核验来源？", "", ""),
        ("□", "待审核文字、图片、视频和原型版本的发送方式与反馈周期。", "", ""),
        ("□", "下次沟通的时间、主题、参与人和需提前准备的材料。", "", ""),
    ], detail_header="补充问题", min_row_mm=11)
    add_heading(doc, "五、55—60分钟收尾话术", 2)
    add_callout(doc, "主持话术｜", "我复述一下今天的结论：首版路线是／暂定为______；三个节点是______；不能失真的内容是______；可使用范围是______；公开前需由______复核。我们会在24小时内发送一份简短确认稿，未确认内容不会进入公开版本。请问这段复述有什么需要纠正？", fill=PALE_GOLD)
    add_form(doc, [
        ("专家纠正／补充", "____________________________________________________________\n____________________________________________________________"),
        ("审核联系人／周期", "联系人：____________　渠道：____________　预计：____个工作日"),
        ("下次沟通", "时间：____________　主题：________________　参与人：____________"),
    ], min_row_mm=13)
    save(doc, FILES[4])


def build_05() -> None:
    title = "拍摄录音及素材使用授权确认单"
    doc = base_doc("05", title)
    add_title(doc, "05", title, "逐项记录本次访谈、人物信息和素材的可用范围；本表不替代历史资料、作品或其他权利人的授权。")
    add_callout(doc, "填写说明｜", "每项只选择一种：□ 同意　□ 不同意　□ 需发布前复核。空白不视为同意；“内部研究”不等于“公开传播”。可随时提出变更，变更前已合法完成的使用按双方确认处理。", fill=SOFT_RED)
    add_form(doc, [("受访人", "姓名：________________　公开称呼：________________"),
                   ("访谈信息", "日期：____年__月__日　地点：________________"),
                   ("项目联系人", "姓名：________________　手机／微信：________________")], min_row_mm=9)
    add_heading(doc, "一、现场记录方式", 2)
    add_table(doc, ["项目", "同意", "不同意", "发布前复核", "限制／说明"], [
        ("现场录音", "□", "□", "□", ""), ("现场录像", "□", "□", "□", ""),
        ("现场拍照", "□", "□", "□", ""), ("书面笔记与原话摘录", "□", "□", "□", ""),
    ], [2.2, 0.75, 0.75, 1.05, 2.75], font_size=8.5, center_cols=(1,2,3), zebra=False, min_row_mm=8)
    add_page_break(doc)
    add_heading(doc, "二、人物信息与素材类型", 2)
    add_table(doc, ["项目", "同意", "不同意", "发布前复核", "限制／说明"], [
        ("人物姓名", "□", "□", "□", ""), ("公开身份", "□", "□", "□", ""),
        ("肖像", "□", "□", "□", ""),
        ("声音", "□", "□", "□", ""), ("工艺过程", "□", "□", "□", ""),
        ("工具与材料", "□", "□", "□", ""), ("访谈／制作场所", "□", "□", "□", ""),
        ("本人拥有权利的作品", "□", "□", "□", "具体作品另列"),
    ], [2.2, 0.75, 0.75, 1.05, 2.75], font_size=8.3, center_cols=(1,2,3), zebra=False, min_row_mm=7.5)

    add_page_break(doc)
    add_heading(doc, "三、具体使用场景（分别授权）", 1)
    add_table(doc, ["使用场景", "同意", "不同意", "发布前复核", "范围／限制"], [
        ("团队内部研究、复盘与原型设计", "□", "□", "□", "仅项目团队／指导教师"),
        ("小程序内部测试", "□", "□", "□", "测试对象：________"),
        ("公开小程序内容与功能展示", "□", "□", "□", "页面／素材：________"),
        ("广东省攀登计划申报材料", "□", "□", "□", "正文／附件：________"),
        ("攀登计划答辩与成果展示", "□", "□", "□", "现场／录播：________"),
        ("校园讲座、工作坊或成果展", "□", "□", "□", "学校／场次：________"),
        ("短视频及其他公开传播", "□", "□", "□", "平台／账号：________"),
        ("论文、报告或研究成果", "□", "□", "□", "匿名／署名：________"),
    ], [2.65, 0.7, 0.7, 1.05, 2.4], font_size=8.35, center_cols=(1,2,3), zebra=False, min_row_mm=10)
    add_callout(doc, "授权隔离｜", "若只同意“团队内部研究”，相关内容不得进入申报附件、答辩、校园公开活动、短视频或公开小程序。", fill=SOFT_RED)
    add_heading(doc, "四、授权条件", 2)
    add_form(doc, [
        ("授权期限", "□ 本次项目阶段：____年__月__日至____年__月__日　□ 其他：____________"),
        ("地域／平台", "____________________________________________________________"),
        ("署名方式", "□ 实名：____________　□ 指定称谓：____________　□ 匿名"),
        ("审核联系人", "姓名：____________　渠道：____________　反馈周期：____个工作日"),
        ("变更／撤回方式", "联系项目联系人并说明具体素材、用途和生效时间；团队书面确认处理结果。"),
        ("其他限制", "____________________________________________________________\n____________________________________________________________"),
    ], min_row_mm=11)

    add_page_break(doc)
    add_heading(doc, "五、历史资料、作品与第三方材料登记", 1)
    add_callout(doc, "权利提示｜", "以下仅登记来源线索，不构成授权。历史照片、作品图片、旧视频、他人肖像、场地标识或第三方作品须另行确认权利归属和使用范围。权属不清时不得复制、发布或用于训练模型。", fill=SOFT_RED)
    add_table(doc, ["材料名称／描述", "来源／持有人", "可能权利人", "拟用途", "后续授权状态"], [
        ("", "", "", "", "□ 待联系"), ("", "", "", "", "□ 待联系"),
        ("", "", "", "", "□ 待联系"),
    ], [1.7, 1.5, 1.45, 1.4, 1.45], font_size=8.4, zebra=False, min_row_mm=17)
    add_heading(doc, "六、确认与签署", 2)
    add_body(doc, "本人已阅读并理解上述逐项选择，知道可对不同记录方式、素材类型和使用场景作出不同决定；未勾选项目不视为同意。", size=9.3)
    add_form(doc, [
        ("受访人签名", "________________　日期：____年__月__日"),
        ("项目方代表签名", "________________　日期：____年__月__日"),
        ("代理人／机构代表", "姓名：____________　身份及权限依据：____________________________"),
        ("如需发布前复核", "□ 已确认审核联系人　□ 已约定发送渠道　□ 已约定反馈周期"),
        ("变更记录", "日期：________　内容：________________________________________"),
    ], min_row_mm=12)
    save(doc, FILES[5])


def build_06() -> None:
    title = "访谈对象身份与联系人信息确认表"
    doc = base_doc("06", title)
    add_title(doc, "06", title, "准确记录身份、公开表述和后续联系人；传承人级别或资格只在有证据时填写。")
    add_heading(doc, "一、基本身份", 2)
    add_form(doc, [
        ("姓名", "________________"), ("常用称呼", "________________"),
        ("所在机构", "____________________________________________________________"),
        ("职务／专业身份", "____________________________________________________________"),
        ("与金箔工艺的关系", "____________________________________________________________"),
        ("建议公开表述", "____________________________________________________________"),
        ("不应使用的称谓", "____________________________________________________________"),
    ], min_row_mm=10)
    add_heading(doc, "二、资格或传承人级别（有证据才填写）", 2)
    add_callout(doc, "核验原则｜", "未看到证书、主管部门公示或可靠公开来源时，保持空白或标为“待核实”，不得根据口头印象推定级别。", fill=SOFT_RED)
    add_table(doc, ["资格／级别", "颁发或公示单位", "证据名称／链接／存档位置", "核验状态"], [
        ("", "", "", "□ 已核验 □ 待核实"), ("", "", "", "□ 已核验 □ 待核实"),
    ], [1.55, 1.7, 2.9, 1.35], font_size=8.5, zebra=False, min_row_mm=16)

    add_page_break(doc)
    add_heading(doc, "三、联系方式与沟通偏好", 1)
    add_form(doc, [
        ("手机", "________________"), ("微信", "________________"), ("邮箱", "________________"),
        ("首选渠道", "□ 微信　□ 电话　□ 邮件　□ 其他：____________"),
        ("适合联系时段", "____________________________________________________________"),
        ("内容审核联系人", "姓名：____________　关系／职务：____________　渠道：____________"),
        ("授权联系人", "姓名：____________　关系／职务：____________　渠道：____________"),
        ("紧急联系人", "姓名：____________　关系：____________　电话：____________"),
    ], min_row_mm=10)
    add_heading(doc, "四、信息公开分级", 2)
    add_table(doc, ["级别", "可记录内容", "具体信息／限制"], [
        ("可公开", "可进入经审核的小程序、申报、答辩或公开传播", ""),
        ("仅内部使用", "仅项目团队和指导教师研究／联络使用", ""),
        ("不得记录／保存", "不进入笔记、录音录像或资料库", ""),
    ], [1.25, 3.25, 3.0], font_size=8.6, center_cols=(0,), zebra=False, min_row_mm=16)
    add_heading(doc, "五、身份信息核验记录", 2)
    add_form(doc, [
        ("核验来源", "□ 本人确认　□ 证书／公示　□ 机构网页　□ 其他：____________"),
        ("来源位置", "____________________________________________________________"),
        ("核验人／日期", "________________／____年__月__日"),
        ("会后是否需本人复核", "□ 是，发送内容：________________　□ 否"),
        ("填写人／确认人", "填写人：____________　信息确认人：____________"),
        ("确认方式／日期", "□ 现场签字 □ 微信 □ 邮件 □ 其他：________　____年__月__日"),
        ("后续变更记录", "日期：________　变更内容：______________________________________"),
    ], min_row_mm=10)
    save(doc, FILES[6])


def build_07() -> None:
    title = "时间地点确认及微信约访文本"
    doc = base_doc("07", title)
    add_title(doc, "07", title, "提供可直接复制的正式约访、确认、提醒和变更文本；发送前替换方括号中的信息。")
    add_form(doc, [("对方称呼", "________________"), ("发送人／身份", "________________"),
                   ("联系电话／微信", "________________"), ("预计参与人数", "____人")], min_row_mm=8)
    add_callout(doc, "授权边界｜", "微信约访与时间地点确认只代表参与意愿，不等同正式授权。录音录像须现场再次确认，并以05授权确认单记录具体使用范围。", fill=SOFT_RED)
    add_heading(doc, "模板1｜首次正式约访", 2)
    add_callout(doc, "可复制文本｜", "[老师称呼]您好！我是[学校／团队][姓名]。我们正在开展金箔非遗数字体验项目，首阶段面向广东高校学生，计划以“3分钟轻量工艺体验＋可选深度学习”为主线。目前制箔、贴箔及具体互动节点都还是候选方案，想请您从专业角度帮助判断真实性、公开边界和适合首版的内容。想预约一次约60分钟的访谈，预计[人数]人参加。我们会带纸面原型；如需录音、录像或拍照，会在现场逐项征得您的同意。如您方便，也请准备可供现场查看的材料线索；权属不明的原文件无需提前发送。请问您近期何时方便？感谢！", fill=PALE_GOLD)
    add_heading(doc, "发送前确认", 3)
    add_checklist(doc, [
        ("□", "学校、团队、发送人身份准确", "", ""), ("□", "未把候选路线写成既定结论", "", ""),
        ("□", "已说明时长、人数、主题、纸面原型和记录需同意", "", ""),
    ], detail_header="检查项", min_row_mm=8)

    add_page_break(doc)
    add_heading(doc, "模板2｜时间地点确认", 1)
    add_callout(doc, "可复制文本｜", "[老师称呼]您好，感谢您接受访谈。我们再次确认：时间为[日期][开始时间—结束时间]，地点／线上方式为[地点或链接]，预计[人数]人参加，主题为金箔数字轻体验的候选路线、三个互动节点、真实性红线及素材使用范围。我们会携带纸面原型和授权确认单；是否录音、录像或拍照均以您的现场选择为准。若需要我们提前准备或避免携带、拍摄某些材料，请随时告诉我。联系人：[姓名＋电话／微信]。")
    add_form(doc, [
        ("日期时间", "____年__月__日　____—____"), ("地点／链接", "____________________________________________________________"),
        ("交通／门禁", "____________________________________________________________"),
        ("对方希望准备的材料", "____________________________________________________________"),
        ("不可携带／不可拍摄", "____________________________________________________________"),
        ("对方确认时间", "____年__月__日　____:____　□ 微信 □ 电话 □ 邮件"),
    ], min_row_mm=10)

    add_page_break(doc)
    add_heading(doc, "模板3｜访谈前一天提醒", 1)
    add_callout(doc, "可复制文本｜", "[老师称呼]您好，提醒您我们明天的访谈安排：时间[开始—结束]，地点[地点／链接]，预计[人数]人参加。我们会用约5—10分钟展示纸面原型，再请您重点帮助确认候选工艺路线、三个互动节点、真实性红线和可使用范围。记录方式仍以您的现场选择为准。如时间、地点或接待要求有变化，请联系[姓名＋电话／微信]。明天见，感谢！")
    add_checklist(doc, [
        ("□", "对方已回复或已读；未回复时用约定渠道再次确认", "", ""),
        ("□", "参与人、交通、门禁、停车／线上链接已确认", "", ""),
        ("□", "设备、打印件和离线备份已按08清单检查", "", ""),
    ], detail_header="提醒后动作", min_row_mm=10)

    add_page_break(doc)
    add_heading(doc, "模板4｜临时变更、取消与重新确认", 1)
    add_callout(doc, "我方申请变更｜", "[老师称呼]您好，很抱歉因[简要原因]需要调整原定于[原时间]的访谈。请问以下时间是否有一项方便：[选项1]／[选项2]／[选项3]？访谈时长、主题和记录授权方式不变。如这些时间均不方便，我们完全配合您的安排。给您带来不便，十分抱歉。")
    add_callout(doc, "对方提出变更后的回复｜", "收到，感谢您告知。我们将原安排标记为取消／待调整。根据您的时间，我们建议改为[新日期][新时间]，地点／方式为[新地点或链接]，预计[人数]人参加。请您确认是否合适；确认后我们会更新全部团队成员和材料。")
    add_form(doc, [
        ("原安排", "____________________________________________________________"),
        ("变更原因（内部）", "____________________________________________________________"),
        ("新安排", "____________________________________________________________"),
        ("对方确认", "____年__月__日　____:____　确认人：____________"),
        ("团队已同步", "□ 全员　□ 设备负责人　□ 交通／场地　□ 文件版本"),
    ], min_row_mm=11)
    save(doc, FILES[7])


def build_08() -> None:
    title = "现场设备与记录工具检查清单"
    doc = base_doc("08", title)
    add_title(doc, "08", title, "在出发前、到场后和离场前逐项检查，降低设备、打印、环境和资料回收失败风险。")
    add_form(doc, [("访谈日期", "____年__月__日"), ("地点", "________________"),
                   ("总检查人", "________________"), ("备用联络人", "________________")], min_row_mm=8)
    add_heading(doc, "一、出发前｜设备与电源", 2)
    add_checklist(doc, [
        ("□", "主录音／录像手机电量≥80%，存储空间充足", "", ""),
        ("□", "已进行30秒录音测试并回放，声音清晰", "", ""),
        ("□", "备用手机／录音设备可用，时间设置正确", "", ""),
        ("□", "充电宝、充电线、转接头齐全且可用", "", ""),
        ("□", "静音设置完成：关闭通知音、闹钟和自动锁屏干扰", "", ""),
        ("□", "拍摄分辨率与文件格式已确认，不使用不必要的超大规格", "", ""),
    ], min_row_mm=8)
    add_page_break(doc)
    add_heading(doc, "二、出发前｜纸质与电子材料", 2)
    add_checklist(doc, [
        ("□", "01—06、08纸质文件建议各2份，页码和版本正确", "", ""),
        ("□", "纸面原型逐屏完整，无已确定路线的误导文字", "", ""),
        ("□", "签字笔≥3支、便签、夹板／文件夹齐全", "", ""),
        ("□", "十份Word与原型PDF已存手机／电脑，并有离线备份", "", ""),
        ("□", "文件名采用日期_对象_类型_版本，原始与公开目录分开", "", ""),
    ], min_row_mm=8)

    add_page_break(doc)
    add_heading(doc, "三、到场后｜环境与座位", 1)
    add_checklist(doc, [
        ("□", "确认光线、逆光、噪声、回声、空调／机器声", "", ""),
        ("□", "确认网络；线上材料已准备离线替代", "", ""),
        ("□", "座位便于交流，设备不遮挡受访人视线", "", ""),
        ("□", "确认拍摄区域、允许拍摄的物品和背景；移除敏感信息", "", ""),
        ("□", "主设备与备用设备分别放置，避免相互摩擦或遮挡", "", ""),
        ("□", "纸面原型、提纲、授权单按使用顺序摆放", "", ""),
    ], min_row_mm=10)
    add_heading(doc, "四、开始记录前｜授权与同步", 2)
    add_checklist(doc, [
        ("□", "已说明项目、时长、记录方式和可拒绝权利", "", ""),
        ("□", "已按05逐项确认录音、录像、拍照；未同意项目不开启", "", ""),
        ("□", "两台设备时间基本一致；记录开始时间点", "", ""),
        ("□", "试录5—10秒并确认声音／画面；记录文件编号", "", ""),
        ("□", "如不允许录音录像，改用两人书面记录并交叉核对", "", ""),
    ], min_row_mm=10)

    add_page_break(doc)
    add_heading(doc, "五、访谈中与离场前", 1)
    add_checklist(doc, [
        ("□", "设备持续工作；每20分钟由设备负责人目视检查一次", "", ""),
        ("□", "重要结论记录时间点；负面意见保留原话", "", ""),
        ("□", "结束前复述路线、三个节点、真实性红线和审核安排", "", ""),
        ("□", "授权单、身份表、纸面原型批注、便签均已回收", "", ""),
        ("□", "确认是否有不得保存或需当场删除的内容并执行", "", ""),
        ("□", "确认会后联系渠道、审核人、反馈周期和下次安排", "", ""),
        ("□", "离场前检查座位、插座、桌面，无设备和资料遗留", "", ""),
    ], min_row_mm=10)
    add_heading(doc, "六、文件回收与首次备份", 2)
    add_form(doc, [
        ("主文件名", "____________________________________________________________"),
        ("备用文件名", "____________________________________________________________"),
        ("首次备份位置", "____________________________________________________________"),
        ("纸质原件保管人", "________________"),
        ("异常与补救", "____________________________________________________________\n____________________________________________________________"),
        ("最终检查", "检查人：____________　时间：____年__月__日 ____:____"),
    ], min_row_mm=11)
    save(doc, FILES[8])


def build_09() -> None:
    title = "访谈纪要关键结论与会后确认单"
    doc = base_doc("09", title)
    add_title(doc, "09", title, "访谈后24小时内形成可复核结论，区分确认状态，并转换为四岗位第2周行动。")
    add_form(doc, [
        ("访谈对象", "________________"), ("日期／地点", "____年__月__日／________________"),
        ("参与人员", "____________________________________________________________"),
        ("主持／记录", "________________／________________"),
        ("记录文件", "录音：____________　视频：____________　照片：____________"),
        ("关联05授权单", "版本／编号：________________　签署日期：____年__月__日"),
    ], min_row_mm=9)
    add_heading(doc, "一、核心观点与原话摘录", 2)
    add_table(doc, ["时间点／来源", "专家原话或明确观点", "适用范围／上下文", "需本人复核"], [
        ("", "", "", "□"), ("", "", "", "□"), ("", "", "", "□"), ("", "", "", "□"),
    ], [1.2, 3.25, 2.3, 0.85], font_size=8.3, center_cols=(0,3), zebra=False, min_row_mm=16)

    add_page_break(doc)
    add_heading(doc, "二、首版路线与三个互动节点", 1)
    add_form(doc, [
        ("路线结论", "□ 制箔候选　□ 贴箔候选　□ 第三种：____________　□ 暂不确定"),
        ("专业理由／适用条件", "____________________________________________________________\n____________________________________________________________"),
        ("结论状态", "□ 已确认　□ 待本人复核　□ 多方意见冲突　□ 需补证据"),
    ], min_row_mm=13)
    add_table(doc, ["节点", "名称／顺序", "数字交互表达", "不可简化／风险", "状态"], [
        ("1", "", "", "", "□确认 □待核"), ("2", "", "", "", "□确认 □待核"),
        ("3", "", "", "", "□确认 □待核"),
    ], [0.55, 1.65, 2.25, 2.35, 1.0], font_size=8.2, center_cols=(0,4), zebra=False, min_row_mm=18)
    add_heading(doc, "三、真实性红线、公开边界与审核", 2)
    add_table(doc, ["类别", "具体内容", "适用范围／例外", "审核人"], [
        ("不得错误表述", "", "", ""), ("不得展示／记录", "", "", ""),
        ("可公开但需审核", "", "", ""), ("可直接用于内部", "", "", ""),
    ], [1.45, 3.0, 2.0, 1.25], font_size=8.4, center_cols=(0,), zebra=False, min_row_mm=15)

    add_page_break(doc)
    add_heading(doc, "四、授权摘要与待补手续", 1)
    add_table(doc, ["内容／用途", "已同意", "不同意", "发布前复核", "待补签／限制"], [
        ("录音／录像／照片", "□", "□", "□", ""), ("姓名／身份／肖像／声音", "□", "□", "□", ""),
        ("工艺／工具材料／场所／作品", "□", "□", "□", ""), ("内部研究", "□", "□", "□", ""),
        ("小程序／申报／答辩", "□", "□", "□", ""), ("校园活动／短视频／公开传播", "□", "□", "□", ""),
    ], [2.25, 0.7, 0.7, 1.1, 2.85], font_size=8.3, center_cols=(1,2,3), zebra=False, min_row_mm=10)
    add_callout(doc, "以签署件为准｜", "本页是摘要，不扩大05授权确认单的范围。历史照片、作品图片和第三方材料须单独确认权利归属。", fill=SOFT_RED)
    add_heading(doc, "五、结论状态清单", 2)
    add_table(doc, ["状态", "事项", "责任人", "期限／证据"], [
        ("已确认", "", "", ""), ("待确认", "", "", ""),
        ("被否定", "", "", ""), ("需补证据", "", "", ""),
    ], [1.05, 3.55, 1.15, 2.05], font_size=8.6, center_cols=(0,2), zebra=False, min_row_mm=15)

    add_page_break(doc)
    add_heading(doc, "六、四岗位第2周行动", 1)
    add_table(doc, ["岗位", "下一步行动", "负责人", "截止时间", "交付证据"], [
        ("岗位1 产品与交互", "把已确认路线、节点和红线转为低保真流程；未确认内容继续标识。", "", "", "原型链接／截图"),
        ("岗位2 技术开发", "验证基础交互和候选3D／AR可行性，不制作未经确认的公开工艺内容。", "", "", "代码／测试记录"),
        ("岗位3 专业对接与传播", "发送确认稿，跟进授权、审核人、下一次资料对接。", "", "", "回复／授权件"),
        ("岗位4 研究与证据", "整理原话、来源、冲突意见和证据缺口，建立研究台账。", "", "", "纪要／证据表"),
    ], [1.45, 3.45, 1.0, 1.0, 1.3], font_size=8.05, center_cols=(0,2,3), zebra=False, min_row_mm=14)
    add_heading(doc, "七、发给专业老师的简短确认稿", 2)
    add_callout(doc, "可复制文本｜", "[老师称呼]您好，感谢今天的指导。我们理解的关键结论如下：①首版路线为／仍待确认[____]；②三个关键节点为[____]；③不得失真的内容为[____]；④允许使用范围为[____]，其中[____]需发布前复核；⑤后续由[联系人]通过[渠道]审核。未确认内容我们不会进入公开版本。烦请您回复“确认”，或直接指出需要修改之处，谢谢！", fill=PALE_GOLD)
    add_form(doc, [
        ("发送时间／渠道", "____年__月__日 ____:____　□ 微信 □ 邮件 □ 其他：________"),
        ("确认人／方式", "________________　□ 明确回复确认 □ 修改后确认 □ 仅收到未确认"),
        ("确认日期", "____年__月__日"),
        ("修改记录", "____________________________________________________________\n____________________________________________________________"),
        ("最终状态", "□ 可进入第2周　□ 部分进入，限制：____________　□ 暂停公开内容"),
    ], min_row_mm=12)
    save(doc, FILES[9])


BUILDERS = [build_00, build_01, build_02, build_03, build_04, build_05, build_06, build_07, build_08, build_09]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for filename in FILES:
        path = OUT_DIR / filename
        if path.exists():
            path.unlink()
    for builder in BUILDERS:
        builder()
    print(f"generated={len(FILES)}")
    for filename in FILES:
        path = OUT_DIR / filename
        print(f"{path.name}\t{path.stat().st_size}")


if __name__ == "__main__":
    main()
