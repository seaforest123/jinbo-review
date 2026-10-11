from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "gold-foil-heritage" / "金箔非遗数字体验与活化传播项目简报V0.2.docx"
TABLE_GEOMETRY_DIR = Path(
    r"C:\Users\Andrew\.codex\plugins\cache\openai-primary-runtime\documents\26.715.12143\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_GEOMETRY_DIR))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


# Resolved design system: standard_business_brief with named Chinese/A4 and gold-foil overrides.
FONT_LATIN = "Microsoft YaHei"
FONT_EAST_ASIA = "微软雅黑"
INK = "332B22"
MUTED = "6F675F"
GOLD = "A5781E"
GOLD_DARK = "6F5016"
CREAM = "F7F1E3"
CREAM_LIGHT = "FBF8F1"
GRID = "D8CDB7"
WHITE = "FFFFFF"
SOFT_GRAY = "F2F2F2"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, color=INK, italic=None):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_style_font(style, size, color, bold=False):
    style.font.name = FONT_LATIN
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color=GRID, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_table_borders(table, color=GRID, size="5"):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, color, size)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_fill(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_left_border(paragraph, color=GOLD, size="20", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def next_num_id(numbering):
    ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    return max(ids, default=0) + 1


def next_abstract_id(numbering):
    ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    return max(ids, default=-1) + 1


def create_numbering(document, kind="decimal"):
    numbering = document.part.numbering_part.element
    abstract_id = next_abstract_id(numbering)
    num_id = next_num_id(numbering)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    # OOXML requires all abstractNum elements to precede num instances.
    # Insert before the first existing w:num instead of appending after it.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_body(doc, text, bold_lead=None, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=10.5, bold=True, color=GOLD_DARK)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=10.5)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5)
    return p


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(label)
    set_run_font(r1, size=10.5, bold=True, color=GOLD_DARK)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    return p


def add_bullet(doc, text, bullet_id):
    p = doc.add_paragraph()
    apply_numbering(p, bullet_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_numbered(doc, text, decimal_id):
    p = doc.add_paragraph()
    apply_numbering(p, decimal_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_callout(doc, label, text, fill=CREAM_LIGHT):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.25
    set_paragraph_fill(p, fill)
    set_paragraph_left_border(p)
    r1 = p.add_run(label)
    set_run_font(r1, size=10.5, bold=True, color=GOLD_DARK)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), size={1: 17, 2: 13.5, 3: 11.5}[level], bold=True,
                 color={1: INK, 2: GOLD_DARK, 3: GOLD_DARK}[level])
    return p


def fill_cell(cell, text, bold=False, color=INK, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, weights, header_fill=GOLD_DARK, zebra=True, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        fill_cell(table.rows[0].cells[idx], header, bold=True, color=WHITE, size=9.3,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        table.rows[0].cells[idx].paragraphs[0].paragraph_format.keep_with_next = True
        set_cell_shading(table.rows[0].cells[idx], header_fill)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if cidx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(row.cells[cidx], str(value), size=font_size, align=align)
            if zebra and ridx % 2 == 1:
                set_cell_shading(row.cells[cidx], CREAM_LIGHT)

    section = doc.sections[-1]
    content_width = section_content_width_dxa(section)
    table_width = content_width - 120
    widths = column_widths_from_weights(weights, table_width)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=table_width,
        indent_dxa=120,
        cell_margins_dxa={"top": 105, "bottom": 105, "start": 120, "end": 120},
    )
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_metric_strip(doc, metrics):
    table = doc.add_table(rows=1, cols=len(metrics))
    row = table.rows[0]
    prevent_row_split(row)
    for idx, (value, label) in enumerate(metrics):
        cell = row.cells[idx]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, CREAM)
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        set_run_font(p1.add_run(value), size=17, bold=True, color=GOLD_DARK)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(label), size=8.8, color=MUTED)
    section = doc.sections[-1]
    content_width = section_content_width_dxa(section)
    table_width = content_width - 120
    widths = column_widths_from_weights([1] * len(metrics), table_width)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=table_width,
        indent_dxa=90,
        cell_margins_dxa={"top": 140, "bottom": 140, "start": 90, "end": 90},
    )
    set_table_borders(table, color="E6D8BA", size="4")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.05)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.5, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 17, INK, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13.5, GOLD_DARK, bold=True)
    h2.paragraph_format.space_before = Pt(11)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 11.5, GOLD_DARK, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("金箔非遗数字体验与活化传播项目"), size=8.5, bold=True, color=MUTED)
    set_run_font(hp.add_run("    |    项目简报 V0.2"), size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("第 "), size=9, color=MUTED)
    add_field(fp, "PAGE")
    set_run_font(fp.add_run(" 页"), size=9, color=MUTED)

    doc.core_properties.title = "金箔非遗数字体验与活化传播项目简报 V0.2"
    doc.core_properties.subject = "广东省攀登计划项目准备工作文件"
    doc.core_properties.author = "项目团队"
    doc.core_properties.keywords = "金箔, 非遗, 广东高校, 3D互动, 项目简报"


def build_document():
    doc = Document()
    configure_document(doc)
    problem_bullet_id = create_numbering(doc, "bullet")
    principle_bullet_id = create_numbering(doc, "bullet")
    collaboration_bullet_id = create_numbering(doc, "bullet")
    selection_decimal_id = create_numbering(doc, "decimal")
    verification_decimal_id = create_numbering(doc, "decimal")
    next_docs_bullet_id = create_numbering(doc, "bullet")

    # Opening block: customer_pack pattern, adapted for a formal project brief.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("广东省攀登计划 · 项目前期工作文件"), size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("金箔非遗数字体验与活化传播项目"), size=26, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    subtitle.paragraph_format.keep_with_next = True
    set_run_font(subtitle.add_run("项目简报 V0.2｜从项目共识到可执行任务"), size=12.5, color=MUTED)

    metadata = add_table(
        doc,
        ["项目要素", "当前定义"],
        [
            ("核心用户", "广东高校学生，重点为18—24岁青年"),
            ("核心目标", "构建并验证一套面向高校青年的金箔非遗数字化传播模式"),
            ("核心产品", "3分钟3D工艺互动＋可选AR展示＋真实内容与数字档案"),
            ("团队结构", "4名学生、4个主责岗位、三条工作路线并行推进"),
        ],
        [1.2, 4.8],
        header_fill=GOLD_DARK,
        zebra=False,
        font_size=9.5,
    )
    for ridx in range(1, len(metadata.rows)):
        set_cell_shading(metadata.rows[ridx].cells[0], CREAM)
        fill_cell(metadata.rows[ridx].cells[0], metadata.rows[ridx].cells[0].text, bold=True,
                  color=GOLD_DARK, size=9.4, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_callout(
        doc,
        "一句话定位｜",
        "面向广东高校青年的金箔非遗轻量化数字体验与活化传播平台，以3D工艺互动为核心，连接传承人故事、数字档案与线下实践。",
    )

    add_heading(doc, "执行摘要", 1)
    add_labeled_paragraph(
        doc,
        "要解决的问题｜",
        "高校青年对金箔非遗认知有限，现有传播多为静态图文或单向展示，缺少低门槛、强互动、可持续参与且能够验证效果的数字体验。",
    )
    add_labeled_paragraph(
        doc,
        "项目方法｜",
        "以3分钟3D贴箔互动吸引首次接触，以传承人真实口述和数字档案建立文化理解，再通过校园活动、问卷、访谈和行为数据检验传播效果。",
    )
    add_labeled_paragraph(
        doc,
        "申报逻辑｜",
        "项目不是单纯开发小程序，而是形成“产品＋内容＋研究＋实践＋证据”的完整成果链。",
    )

    add_metric_strip(
        doc,
        [("800+", "真实用户目标"), ("300+", "有效问卷"), ("3场", "校园实践"), ("8—12期", "系列短视频")],
    )

    add_heading(doc, "1. 项目定义", 1)
    add_heading(doc, "1.1 核心目标与价值", 2)
    add_body(
        doc,
        "本项目的核心目标是构建并验证一套面向广东高校青年的金箔非遗数字化传播模式。小程序是体验工具，数字档案是真实内容基础，传承人内容是传播载体，校园活动和用户研究用于验证效果。",
    )
    add_body(
        doc,
        "最终成果由产品、内容、研究、知识产权、社会实践和证据材料共同组成。项目是否成功，不以“是否上线”单独判断，而以用户是否愿意进入、能否完成体验、认知是否提升以及是否产生进一步参与行为综合判断。",
    )

    add_heading(doc, "1.2 目标用户", 2)
    add_labeled_paragraph(
        doc,
        "第一阶段核心用户｜",
        "广东高校学生，重点年龄18—24岁。该群体习惯短视频和轻量互动，对非遗可能有好奇心，但通常缺少主动学习意愿和长时间阅读耐心。",
    )
    add_labeled_paragraph(
        doc,
        "延伸用户｜",
        "高校教师和校园文化活动组织者、非遗爱好者、青年文旅用户、传承人及非遗保护与传播相关机构。产品不设置排他性年龄限制。",
    )

    add_heading(doc, "1.3 项目问题链", 2)
    for item in [
        "传播入口偏重知识灌输，难以吸引青年完成首次接触。",
        "照片、视频、口述和工艺资料分散，缺少统一整理和规范标注。",
        "部分数字化项目重视觉技术、轻真实人物和工艺依据。",
        "传播结果多停留在浏览量，缺少前后测、行为数据和线下实践证据。",
    ]:
        add_bullet(doc, item, problem_bullet_id)

    add_heading(doc, "2. 解决方案与用户路径", 1)
    add_heading(doc, "2.1 “一核三支撑”结构", 2)
    add_labeled_paragraph(doc, "核心体验｜", "建设一套代表性“3分钟金箔轻量工艺体验”，以3D贴箔互动作为首版核心入口，并提供可选AR展示或合影功能。")
    add_labeled_paragraph(doc, "内容支撑｜", "采集并整理传承人口述、真实工艺影像、工具、材料、步骤、作品及经审核的文字解说和元数据。")
    add_labeled_paragraph(doc, "传播支撑｜", "形成8—12期系列短视频和具有连续识别度的传承人故事，负责吸引、解释和引导回流。")
    add_labeled_paragraph(doc, "实践验证｜", "通过高校体验活动、问卷、访谈、行为数据和前后测，验证认知、兴趣和参与意愿的变化。")

    add_callout(
        doc,
        "用户路径｜",
        "接触吸引 → 3分钟互动体验 → 了解传承人与真实工艺 → 浏览数字档案 → 参与线下活动或主动传播 → 数据反馈与产品迭代",
    )

    add_heading(doc, "2.2 产品与内容原则", 2)
    for item in [
        "互动优先：首屏不以课程学习和知识测评作为主要入口，先让用户愿意开始。",
        "真实优先：所有工艺步骤、解说文字和数字模拟均需传承人或专业老师审核。",
        "核心可降级：3D主线可独立完成；设备不支持AR或用户拒绝相机权限时，体验仍然完整。",
        "范围可控制：首版只做一个非遗项目、一位核心合作对象、一套代表性互动和一个可选AR场景。",
        "效果可验证：功能、内容和活动均需对应明确数据或证据，不为展示而堆叠功能。",
    ]:
        add_bullet(doc, item, principle_bullet_id)

    add_heading(doc, "3. 创新主张", 1)
    add_labeled_paragraph(doc, "体验方式创新｜", "将传统图文科普转化为3分钟轻量工艺体验，降低青年首次接触非遗的时间和理解门槛。")
    add_labeled_paragraph(doc, "文化叙事创新｜", "把3D互动、传承人真实口述和数字档案串联，使数字操作对应真实人物、材料、工具和工艺流程。")
    add_labeled_paragraph(doc, "传播验证创新｜", "用校园实践、用户行为数据、前后测和访谈检验实际效果，而不是只统计浏览量和点赞量。")
    add_callout(doc, "创新概括｜", "以互动吸引青年，以真实内容建立文化认同，以实践数据验证传播效果。3D/AR是实现目标的手段，不是项目全部创新依据。")

    add_heading(doc, "4. 项目范围与成果包", 1)
    add_heading(doc, "4.1 首期范围", 2)
    add_table(
        doc,
        ["纳入首期", "暂不纳入首期"],
        [
            ("一个金箔非遗项目及一位核心合作对象", "同时覆盖多个非遗门类"),
            ("一个微信小程序", "独立APP、商城、支付和复杂会员体系"),
            ("一套3D贴箔互动＋一个可选AR场景", "大型3D游戏、多关卡重度游戏化和完整空间AR"),
            ("数字档案、传承人故事和系列短视频", "直播带货、商业经纪等重运营业务"),
            ("高校调研、用户测试和校园实践", "为增加数量而申请缺乏创新基础的专利"),
        ],
        [1, 1],
        font_size=9.4,
    )
    add_callout(doc, "范围控制规则｜", "新增功能必须说明其对核心目标、评价指标和申报成果的直接贡献，否则不进入首期交付范围。", fill=SOFT_GRAY)

    add_heading(doc, "4.2 预期成果包", 2)
    add_table(
        doc,
        ["类别", "核心交付", "目标规模"],
        [
            ("产品技术", "微信小程序、3D互动、可选AR、数字档案功能及技术文档", "1套完整可运行成果"),
            ("知识产权", "软件著作权；专利按真实创新条件启动", "软著1项，专利为加分项"),
            ("内容传播", "规范化数字档案、口述与工艺影像专题、系列短视频", "视频8—12期"),
            ("研究评价", "调研报告、论文初稿或投稿成果、问卷与前后测", "问卷300+，前后测100+"),
            ("实践影响", "校园活动、真实用户、合作证明、数据和社会影响材料", "活动3场，用户800+"),
        ],
        [1.1, 3.5, 1.4],
        font_size=9.1,
    )

    add_heading(doc, "5. 效果评价与证据链", 1)
    add_labeled_paragraph(doc, "触达｜", "真实用户不少于800人；校园体验或推广活动不少于3场。")
    add_labeled_paragraph(doc, "体验｜", "进入核心体验后的完成率目标不低于60%；体验后继续查看故事或数字档案的比例目标不低于25%。")
    add_labeled_paragraph(doc, "认知｜", "不少于100名用户完成体验前后测；金箔相关知识得分平均提升目标为20%以上。")
    add_labeled_paragraph(doc, "意愿｜", "比较体验前后用户继续了解、关注、分享或参加线下活动的意愿变化，并结合访谈解释原因。")
    add_labeled_paragraph(doc, "成果｜", "按计划完成小程序、数字档案、系列视频、软件著作权、调研报告、论文初稿和项目证据包。")
    add_callout(doc, "数据原则｜", "以上均为待验证的目标值，不作为已经取得的成绩表述。申报前30—50人的小范围测试用于校准目标，不为满足指标而筛选或修改数据。")

    add_heading(doc, "6. 四岗位团队与协作", 1)
    add_table(
        doc,
        ["岗位", "主责内容", "核心交付", "固定协作"],
        [
            ("1 产品与前端", "用户流程、原型、小程序前端、体验优化", "PRD、页面原型、前端功能、使用说明", "与岗位2共同开发；与岗位3完成内容上架"),
            ("2 3D与技术", "架构、3D接入、互动逻辑、AR、数据与部署", "3D互动、AR场景、接口、技术与兼容文档", "与岗位1维护代码；与岗位4设计埋点"),
            ("3 内容与档案", "合作对接、采访拍摄、资料整理、元数据和视频", "内容手册、数字档案、8—12期视频、授权记录", "与岗位1上架内容；与岗位4开展访谈活动"),
            ("4 调研与实践", "问卷、测试、活动、传播、数据分析和证据归档", "调研方案、300份问卷、活动材料、效果报告", "与岗位2统计数据；与岗位3开展校园实践"),
        ],
        [1.05, 1.75, 1.75, 1.45],
        font_size=8.7,
    )

    add_heading(doc, "6.1 协作规则", 2)
    for item in [
        "岗位1和岗位2共同开发，但分别主责用户可见部分与3D、数据和运行部分。",
        "岗位3提供经过整理和审核的内容，岗位1负责产品化呈现，避免开发人员临时改写文化内容。",
        "岗位4先定义评价指标，岗位2再实施数据埋点，避免产品完成后才补统计。",
        "全体成员参加关键访谈、阶段评审、用户测试和大型活动。",
        "主责岗位意味着对交付结果负责，不意味着成员只能完成该岗位工作。",
    ]:
        add_bullet(doc, item, collaboration_bullet_id)

    add_heading(doc, "6.2 学生选岗方式", 2)
    for item in [
        "每名学生填写第一志愿和第二志愿。",
        "每个岗位安排一个2—3小时的小型试做任务。",
        "结合兴趣、试做结果和团队组合确定最终岗位，并设置交接备份人。",
    ]:
        add_numbered(doc, item, selection_decimal_id)

    add_heading(doc, "7. 双阶段实施计划", 1)
    add_heading(doc, "7.1 申报前准备期：倒排8周", 2)
    add_table(
        doc,
        ["时间", "重点任务", "阶段交付"],
        [
            ("第1—2周", "核实非遗背景、首次访谈、合作意向和素材授权", "身份依据、访谈纪要、合作与授权文件"),
            ("第3—4周", "完成项目简报、PRD、内容手册和技术验证", "三份工作文件、技术可行性结论"),
            ("第5—6周", "制作交互原型、1个3D演示和2期样片", "可演示原型、3D样例、视频样片"),
            ("第7周", "组织30—50名高校学生小范围测试", "测试数据、问题清单、用户反馈"),
            ("第8周", "迭代并整理申报支撑材料", "合作证明、截图、数据和申报证据包"),
        ],
        [1.0, 3.0, 2.0],
        font_size=9.1,
    )
    add_callout(doc, "申报最低基础｜", "真实合作对象＋明确需求＋可演示原型＋样片内容＋首轮用户数据。")

    add_heading(doc, "7.2 立项实施期：12—18个月", 2)
    add_table(
        doc,
        ["阶段", "重点任务", "主要成果"],
        [
            ("第1—3月", "正式资料采集与小程序V1.0", "素材库、访谈资料、产品V1.0"),
            ("第4—6月", "完成3D互动、数字档案与软著申请", "核心体验、档案系统、软著材料"),
            ("第7—10月", "校园活动、短视频传播和规模化测试", "活动证据、传播数据、用户样本"),
            ("第11—14月", "数据驱动迭代、报告和论文整理", "产品迭代、调研报告、论文初稿"),
            ("第15—18月", "成果整合与结项", "知识产权、社会影响和结项证据包"),
        ],
        [1.0, 3.0, 2.0],
        font_size=9.1,
    )

    add_heading(doc, "8. 风险与治理要求", 1)
    add_table(
        doc,
        ["风险", "控制措施", "必须保留的证据"],
        [
            ("身份与资料真实性", "申报前核实项目正式名称、级别、名录和合作对象身份", "官方名录依据、合作证明"),
            ("文化表达失真", "工艺、文字和3D模拟由传承人或专业老师审核；标注为简化体验", "审核记录、版本记录"),
            ("版权与肖像", "采集前取得书面授权；建立音乐、字体、图片、模型和AI素材来源清单", "授权书、许可与来源表"),
            ("3D/AR技术", "3D主线独立可用；AR失败时自动降级；开展真机兼容测试", "兼容清单、测试记录"),
            ("学生与进度", "模块主责、周任务、阶段验收、代码备份和文档交接", "任务表、验收记录、仓库记录"),
            ("调研与隐私", "知情同意、匿名收集、最少必要信息、限制原始数据访问", "同意说明、匿名数据与权限记录"),
        ],
        [1.25, 3.25, 1.5],
        font_size=8.9,
    )

    add_heading(doc, "9. 申报前待核验事项与下一步", 1)
    add_body(
        doc,
        "现阶段可继续进行项目设计，但以下事项必须在申报前准备期第1—2周核实。在完成核验前，对外统一使用“金箔类非遗项目及其传承人或专业合作对象”表述，不擅自填写认定级别和称号。",
    )
    for item in [
        "确认非遗项目正式名称、所在城市、保护单位、认定级别及官方名录来源。",
        "确认传承人或合作老师姓名、身份及可公开介绍内容。",
        "明确可数字化展示的代表性工艺、工具、作品和故事范围。",
        "取得采访、拍摄、数字化展示及申报材料使用授权。",
        "确认合作方可参与的访谈、专业审核、线下活动和答辩支持方式。",
    ]:
        add_numbered(doc, item, verification_decimal_id)

    add_heading(doc, "下一批执行文件", 2)
    for item in [
        "《金箔非遗小程序产品需求文档（PRD）V0.1》",
        "《金箔非遗内容采集与制作手册V0.1》",
        "《四岗位申报前8周任务与验收表》",
        "《用户研究与效果评价方案》",
        "《素材授权、内容审核与证据归档模板包》",
    ]:
        add_bullet(doc, item, next_docs_bullet_id)

    add_callout(doc, "当前决策｜", "项目简报确认后，下一步先编制PRD V0.1，再同步建立内容采集与制作手册。", fill=CREAM)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
