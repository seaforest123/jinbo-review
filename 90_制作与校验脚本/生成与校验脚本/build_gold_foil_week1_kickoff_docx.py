from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_gold_foil_project_brief_docx import (  # noqa: E402
    CREAM,
    CREAM_LIGHT,
    GOLD,
    GOLD_DARK,
    GRID,
    INK,
    MUTED,
    SOFT_GRAY,
    WHITE,
    add_field,
    fill_cell,
    prevent_row_split,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_style_font,
    set_table_borders,
)

TABLE_GEOMETRY_DIR = Path(
    r"C:\Users\Andrew\.codex\plugins\cache\openai-primary-runtime\documents\26.715.12143\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_GEOMETRY_DIR))
from table_geometry import (  # noqa: E402
    apply_table_geometry,
    column_widths_from_weights,
    section_content_width_dxa,
)


OUT = ROOT / "docs" / "gold-foil-heritage" / "金箔非遗项目第1周启动包V0.1.docx"


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, bottom=90, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_min_height(row, mm: float) -> None:
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.2, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 17, INK, bold=True)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(7)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13, GOLD_DARK, bold=True)
    h2.paragraph_format.space_before = Pt(7)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 11.5, GOLD_DARK, bold=True)
    h3.paragraph_format.space_before = Pt(5)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("金箔非遗数字体验与活化传播项目"), size=8.2, bold=True, color=MUTED)
    set_run_font(header.add_run("    |    第1周项目启动包 V0.1"), size=8.2, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("第 "), size=8.5, color=MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.add_run(" 页"), size=8.5, color=MUTED)

    doc.core_properties.title = "金箔非遗项目第1周启动包 V0.1"
    doc.core_properties.subject = "项目启动会、岗位选择、专业访谈、拍摄授权与第1周验收"
    doc.core_properties.author = "项目团队"
    doc.core_properties.keywords = "金箔, 非遗, 攀登计划, 启动包, 第一周, 访谈, 授权, 验收"


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    size = {1: 17, 2: 13, 3: 11.5}[level]
    set_run_font(run, size=size, bold=True, color=INK if level == 1 else GOLD_DARK)
    return paragraph


def add_body(doc: Document, text: str, *, size=10.2, bold=False, color=INK, after=5, align=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.22
    if align is not None:
        paragraph.alignment = align
    set_run_font(paragraph.add_run(text), size=size, bold=bold, color=color)
    return paragraph


def add_callout(doc: Document, label: str, text: str, *, fill=CREAM_LIGHT):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(6)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(paragraph, fill)
    set_run_font(paragraph.add_run(label), size=10.2, bold=True, color=GOLD_DARK)
    set_run_font(paragraph.add_run(text), size=10.2, color=INK)
    return paragraph


def remove_trailing_blank_paragraph(doc: Document) -> None:
    """Remove helper spacing that would otherwise become a blank page."""
    if not doc.paragraphs:
        return
    paragraph = doc.paragraphs[-1]
    if paragraph.text.strip():
        return
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def add_page_title(doc: Document, title: str, subtitle: str, *, break_before=True) -> None:
    if break_before:
        remove_trailing_blank_paragraph(doc)
    heading = add_heading(doc, title, 1)
    heading.paragraph_format.page_break_before = break_before
    add_body(doc, subtitle, size=9.5, color=MUTED, after=8)


def add_table(
    doc: Document,
    headers,
    rows,
    weights,
    *,
    font_size=8.6,
    center_cols=(),
    zebra=True,
    min_row_mm=None,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        fill_cell(
            table.rows[0].cells[idx],
            header,
            bold=True,
            color=WHITE,
            size=8.7,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[idx], GOLD_DARK)
        table.rows[0].cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if min_row_mm is not None:
            set_row_min_height(row, min_row_mm)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            align = WD_ALIGN_PARAGRAPH.CENTER if cidx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(cell, str(value), size=font_size, align=align, color=INK)
            if zebra and ridx % 2 == 1:
                set_cell_shading(cell, CREAM_LIGHT)

    section = doc.sections[-1]
    content_width = section_content_width_dxa(section)
    table_width = content_width - 120
    widths = column_widths_from_weights(weights, table_width)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=table_width,
        indent_dxa=120,
        cell_margins_dxa={"top": 90, "bottom": 90, "start": 120, "end": 120},
    )
    set_table_borders(table, color=GRID, size="5")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_form_table(doc: Document, rows, *, weights=(1.7, 5.3), min_row_mm=9, font_size=9.1):
    return add_table(
        doc,
        ["字段", "填写内容"],
        rows,
        list(weights),
        font_size=font_size,
        center_cols=(0,),
        zebra=False,
        min_row_mm=min_row_mm,
    )


def add_cover(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(3)
    set_run_font(kicker.add_run("广东省攀登计划 · 学生项目执行文件"), size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("金箔非遗项目第1周启动包"), size=26, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    set_run_font(subtitle.add_run("V0.1｜从团队认岗到首次专业访谈与周验收"), size=12, color=MUTED)

    add_table(
        doc,
        ["周期", "团队", "启动会", "共同闸门"],
        [("7月17日—7月23日", "4名学生＋指导教师", "45分钟", "首次专业访谈与工艺主线确认")],
        [1.6, 1.65, 1.2, 2.55],
        font_size=9.0,
        center_cols=(0, 1, 2, 3),
        zebra=False,
        min_row_mm=11,
    )

    add_callout(
        doc,
        "使用原则｜",
        "四个岗位并行推进，但共享同一专业确认闸门。访谈前允许建立内部临时技术占位路线，仅用于交互与性能验证，不得作为已确认工艺或公开文化内容。",
        fill=CREAM,
    )

    add_heading(doc, "本包如何使用", 2)
    add_table(
        doc,
        ["顺序", "表单", "使用时点", "归档结果"],
        [
            ("1", "项目启动会议程", "启动会前后", "会议决定与会后动作"),
            ("2", "岗位志愿与团队分工", "启动会现场", "4个主责岗位与备份协作"),
            ("3", "首次专业访谈准备与提纲", "访谈前／访谈中", "访谈纪要与专业确认"),
            ("4", "拍摄与授权确认清单", "访谈及拍摄准备", "授权边界与素材台账"),
            ("5", "第1周任务与验收记录", "7月23日前后", "通过、返工或阻塞结论"),
        ],
        [0.65, 2.2, 1.5, 2.65],
        font_size=8.7,
        center_cols=(0, 2),
    )

    add_heading(doc, "第1周结束条件", 2)
    add_table(
        doc,
        ["检查", "完成条件"],
        [
            ("□", "4名学生完成岗位志愿，主责人与备份协作者已经确定。"),
            ("□", "首次专业访谈已完成，或已有明确到具体日期与对象的访谈安排。"),
            ("□", "制箔／贴箔首版主线有专业依据；未确认时明确标记为内部临时路线。"),
            ("□", "四个岗位均提交可检查交付物及证据位置。"),
            ("□", "指导教师完成周验收并记录第2周进入条件。"),
        ],
        [0.55, 6.45],
        font_size=9.0,
        center_cols=(0,),
        zebra=False,
    )


def add_agenda(doc: Document) -> None:
    add_page_title(doc, "01｜项目启动会议程", "建议时长45分钟；会前将PRD V0.3和四岗位7周任务表发给全部学生。")
    add_form_table(
        doc,
        [
            ("会议日期", ""),
            ("会议地点／方式", ""),
            ("主持人", "指导教师"),
            ("参会人", "指导教师＋4名学生"),
        ],
        min_row_mm=8,
    )

    add_heading(doc, "45分钟议程", 2)
    add_table(
        doc,
        ["时间", "环节", "主持动作", "必须形成的结果"],
        [
            ("0—5分钟", "项目目标与边界", "说明攀登计划目标、P0范围和第1周共同闸门。", "全员理解首版不做商城、复杂会员或重度游戏。"),
            ("5—12分钟", "核心用户与体验主线", "说明广东高校学生、3分钟互动和可选深度学习。", "确认所有岗位围绕同一核心流程协作。"),
            ("12—20分钟", "四岗位说明", "逐一说明主责、固定协作、交付物和验收人。", "学生明确第一、第二岗位志愿。"),
            ("20—28分钟", "专业访谈闸门", "说明制箔／贴箔待确认、真实性和授权否决条件。", "确定访谈联系人、拟定时间和准备责任。"),
            ("28—38分钟", "个人志愿与分工", "学生填写志愿，教师综合能力与协作关系定岗。", "形成4个主责岗位及备份协作者。"),
            ("38—45分钟", "第1周验收与归档", "确认提交位置、命名规则、周验收时间和状态口径。", "每人复述本周交付物与期限。"),
        ],
        [1.1, 1.35, 2.65, 2.9],
        font_size=8.2,
        center_cols=(0, 1),
    )

    add_heading(doc, "会前检查", 2)
    add_table(
        doc,
        ["检查", "材料／动作", "负责人", "状态"],
        [
            ("□", "PRD V0.3与四岗位7周任务表已发送。", "指导教师", "未开始／已完成"),
            ("□", "4名学生已预读岗位边界并准备志愿理由。", "全体学生", "未开始／已完成"),
            ("□", "首次专业访谈联系人与可选时间已初步确认。", "岗位3候选人", "未开始／已完成"),
            ("□", "共享文件夹、代码仓库和提交命名规则已建立。", "岗位2候选人", "未开始／已完成"),
        ],
        [0.55, 4.0, 1.25, 1.2],
        font_size=8.7,
        center_cols=(0, 2, 3),
        zebra=False,
    )


def add_role_selection(doc: Document) -> None:
    add_page_title(doc, "02｜岗位志愿与团队分工", "每名学生填写第一、第二志愿；指导教师综合能力、兴趣和协作接口确定主责岗位。")
    add_callout(
        doc,
        "定岗规则｜",
        "主责人对最终交付负责，备份协作者必须能够接手关键材料。岗位不是封闭分组：首次访谈、P0评审、用户测试和最终演练由全体参加。",
    )

    add_heading(doc, "个人岗位志愿表", 2)
    add_table(
        doc,
        ["学生姓名", "第一志愿", "选择理由／已有能力", "第二志愿", "可承担的备份事项", "确认"],
        [
            ("", "", "", "", "", ""),
            ("", "", "", "", "", ""),
            ("", "", "", "", "", ""),
            ("", "", "", "", "", ""),
        ],
        [0.85, 1.05, 1.8, 1.05, 1.65, 0.6],
        font_size=8.2,
        center_cols=(0, 1, 3, 5),
        zebra=False,
        min_row_mm=14,
    )
    add_body(doc, "岗位选项：1 产品与前端｜2 3D与技术｜3 内容与档案｜4 调研与实践", size=9.2, color=MUTED, after=6)

    add_heading(doc, "团队定岗汇总", 2)
    add_table(
        doc,
        ["岗位", "主责学生", "备份协作者", "第1周核心交付", "验收人"],
        [
            ("1 产品与前端", "", "", "P0页面清单、待决策清单、产品决策记录", "指导教师＋岗位2"),
            ("2 3D与技术", "", "", "工程骨架、XR基础场景、设备与技术可行性记录", "指导教师＋岗位1"),
            ("3 内容与档案", "", "", "访谈纪要、授权事项清单、首批素材目录", "指导教师"),
            ("4 调研与实践", "", "", "研究目标、样本口径、招募计划、前后测构念", "指导教师"),
        ],
        [1.25, 1.0, 1.05, 2.65, 1.05],
        font_size=8.2,
        center_cols=(0, 1, 2, 4),
        min_row_mm=11,
    )
    add_form_table(doc, [("指导教师确认", "签名／日期：")], min_row_mm=10)


def add_interview_prep(doc: Document) -> None:
    add_page_title(doc, "03｜首次专业访谈准备", "访谈对象可为非遗传承人或能够对工艺真实性负责的专业老师；身份与合作信息后续按证据归档。")
    add_form_table(
        doc,
        [
            ("访谈对象", "姓名／称谓／机构："),
            ("访谈时间", "日期／开始—结束："),
            ("访谈地点／方式", ""),
            ("项目联系人", "姓名／联系方式："),
            ("主访谈人", "建议由岗位3主责；记录人："),
            ("参加人员", "指导教师＋4名学生／实际参加人员："),
            ("记录方式", "□书面纪要　□录音　□录像　□会后书面确认"),
        ],
        min_row_mm=8.5,
    )

    add_heading(doc, "建议开场说明", 2)
    add_callout(
        doc,
        "开场话术｜",
        "我们正在准备一个面向广东高校学生的金箔非遗轻量数字体验项目，计划以3分钟互动模拟连接真实人物、真实工艺影像和后续校园活动。本次访谈用于确认首版工艺主线、真实性边界、拍摄与使用范围。未经明确同意的内容不会进入公开版本；数字互动会标注为简化模拟。",
        fill=CREAM,
    )

    add_heading(doc, "访谈前检查", 2)
    add_table(
        doc,
        ["检查", "确认事项", "负责人", "结果／备注"],
        [
            ("□", "已发送简短项目介绍和预计访谈时长。", "岗位3", ""),
            ("□", "已说明是否希望录音、录像或拍照，并等待对方确认。", "岗位3", ""),
            ("□", "岗位1准备用户流程、3分钟体验和待决策问题。", "岗位1", ""),
            ("□", "岗位2准备候选动作、模型、手势、性能和降级问题。", "岗位2", ""),
            ("□", "岗位4准备研究、活动和受众测试相关问题。", "岗位4", ""),
            ("□", "准备纸质或电子纪要模板，并明确资料保存位置。", "岗位3", ""),
            ("□", "不预设对方必须认可制箔或贴箔路线。", "全体", ""),
        ],
        [0.55, 4.2, 1.0, 1.25],
        font_size=8.5,
        center_cols=(0, 2),
        zebra=False,
        min_row_mm=8,
    )


INTERVIEW_PART_1 = [
    ("1", "对象与身份", "请您介绍自己目前从事的金箔相关工作，以及与该工艺的学习、实践或传承关系。", ""),
    ("2", "对象与身份", "项目在公开材料中可以如何准确称呼您？哪些身份、级别或机构信息需要另行核实？", ""),
    ("3", "对象与身份", "您认为公众对金箔工艺最常见的误解是什么？", ""),
    ("4", "主线选择", "如果首版只能选择一条3分钟体验主线，制箔和贴箔哪一条更有代表性、可拍摄且适合互动？", ""),
    ("5", "主线选择", "您推荐该路线的专业理由是什么？它最能体现哪些材料、工具、动作或价值？", ""),
    ("6", "主线选择", "另一条路线为什么暂不适合首版？后续是否可以作为扩展内容？", ""),
    ("7", "主线选择", "首版路线中有哪些步骤可以简化，哪些步骤即使在数字体验中也不能省略？", ""),
    ("8", "三步互动", "若压缩为三个互动步骤，您建议分别选择哪三个关键节点？", ""),
    ("9", "三步互动", "每个节点应出现哪些工具、材料和动作？正确顺序是什么？", ""),
    ("10", "三步互动", "每个动作最容易被错误表现的地方是什么？", ""),
    ("11", "三步互动", "哪些动作适合拖动、滑动、点击或手势模拟？哪些只适合用真实视频呈现？", ""),
    ("12", "三步互动", "用户完成每一步后，怎样的文字反馈既准确又容易理解？", ""),
    ("13", "三步互动", "结果页可以生成怎样的个人作品或文化提示，才不会误导用户以为完成了真实工艺？", ""),
]


INTERVIEW_PART_2 = [
    ("14", "真实性红线", "哪些术语、人物故事、历史信息或工艺说法必须有来源并经审核？", ""),
    ("15", "真实性红线", "数字互动中应明确标注哪些内容为简化模拟？建议采用怎样的提示语？", ""),
    ("16", "真实性红线", "是否存在不宜拍摄、公开、简化或商业化表达的环节？", ""),
    ("17", "真实性红线", "对画面、服饰、工具摆放、声音和动作节奏有什么专业要求？", ""),
    ("18", "拍摄与使用", "哪些人物、场地、工具、材料、过程和作品可以拍摄？", ""),
    ("19", "拍摄与使用", "照片、视频、声音和文字纪要分别可以用于小程序、申报材料、答辩、短视频和校园展示吗？", ""),
    ("20", "拍摄与使用", "哪些素材需要单独书面授权、署名或发布前再次确认？", ""),
    ("21", "拍摄与使用", "是否可以根据真实工具、材料和动作制作简化3D模型？模型发布前由谁审核？", ""),
    ("22", "拍摄与使用", "已有历史照片、作品图或宣传资料能否使用？其权利人和来源如何确认？", ""),
    ("23", "审核机制", "后续视频、术语、三步脚本、3D动作和人物介绍由谁审核最合适？", ""),
    ("24", "审核机制", "建议每次审核以什么形式提交？正常反馈周期大约多久？", ""),
    ("25", "合作与影响", "除了小程序，您认为校园讲座、工作坊、短视频或文创展示中哪种方式最有传播价值？", ""),
    ("26", "合作与影响", "您希望项目最终帮助解决什么现实问题？哪些成果最能说明项目产生了价值？", ""),
]


def add_interview_questions(doc: Document, title: str, subtitle: str, rows) -> None:
    add_page_title(doc, title, subtitle)
    add_table(
        doc,
        ["序号", "主题", "建议问题", "现场记录关键词"],
        rows,
        [0.55, 1.2, 4.15, 1.1],
        font_size=8.25,
        center_cols=(0, 1),
        zebra=True,
        min_row_mm=8.2,
    )


def add_interview_conclusion(doc: Document) -> None:
    add_page_title(
        doc,
        "05A｜访谈结论与会后确认",
        "访谈结束时先现场复述关键结论；无法当场确认的内容列入会后书面复核，不凭记忆直接进入公开版本。",
    )
    add_form_table(
        doc,
        [
            ("首版工艺主线", "□制箔　□贴箔　□仍需补充证据；专业理由："),
            ("三步框架", "步骤1：　　　　　　步骤2：　　　　　　步骤3："),
            ("不可简化／不可公开", ""),
            ("内容审核人", "姓名／称谓／建议反馈周期："),
            ("确认方式", "□现场复述确认　□会后微信／邮件确认　□书面审核"),
            ("下一次沟通", "预计日期／需补材料："),
        ],
        min_row_mm=10,
        font_size=8.8,
    )
    add_heading(doc, "访谈纪要摘要", 2)
    add_form_table(
        doc,
        [
            ("核心观点", ""),
            ("真实性红线", ""),
            ("对项目的期待", ""),
        ],
        min_row_mm=15,
        font_size=8.8,
    )
    add_heading(doc, "会后确认待办", 2)
    add_table(
        doc,
        ["编号", "待确认内容", "发送人", "确认人", "期限／状态"],
        [(str(i), "", "", "", "未开始") for i in range(1, 5)],
        [0.5, 3.1, 1.0, 1.0, 1.4],
        font_size=8.2,
        center_cols=(0, 2, 3, 4),
        zebra=False,
        min_row_mm=9,
    )


def add_authorization(doc: Document) -> None:
    add_page_title(doc, "06｜拍摄与授权确认清单", "本页用于内部确认和后续授权跟踪，不替代正式肖像、著作权、场地或合作授权文件。")
    add_callout(
        doc,
        "公开底线｜",
        "未授权素材只允许内部联调和评审，不进入公开小程序、短视频、申报附件或答辩展示。专业老师的口头同意应在发布前转化为可归档的书面或电子确认。",
        fill=CREAM,
    )
    add_table(
        doc,
        ["素材／对象", "计划拍摄或制作内容", "拟使用场景", "确认状态", "证据／备注"],
        [
            ("人物肖像", "正面、侧面、操作过程、合影", "小程序／申报／答辩／短视频", "□可用 □另授权 □不可用", ""),
            ("访谈声音", "原声、配音、引用片段", "开场视频／故事页／答辩", "□可用 □另授权 □不可用", ""),
            ("工艺过程", "首版三步及相关真实操作", "步骤视频／内容档案", "□可用 □另授权 □不可用", ""),
            ("工具与材料", "工具外观、材料细节、摆放方式", "3D模型／科普内容", "□可用 □另授权 □不可用", ""),
            ("场地环境", "工作室、展陈、门牌及环境声", "背景画面／项目记录", "□可用 □另授权 □不可用", ""),
            ("作品与应用", "成品、局部细节、应用案例", "结果页／金箔志／海报", "□可用 □另授权 □不可用", ""),
            ("历史资料", "旧照片、证书、图录、报道", "人物故事／申报证据", "□可用 □另授权 □不可用", ""),
            ("姓名与称谓", "姓名、机构、传承身份或专业称谓", "所有公开材料", "□可用 □需核实 □不可用", ""),
            ("3D参考", "依据真实工具、材料和动作制作模型", "小程序互动／AR", "□可用 □审核后用 □不可用", ""),
            ("校园活动", "讲座、体验、工作坊或展示记录", "活动传播／实践证据", "□可用 □另授权 □不可用", ""),
        ],
        [1.15, 2.0, 1.55, 1.35, 1.0],
        font_size=7.8,
        center_cols=(0, 3),
        min_row_mm=8,
    )
    add_heading(doc, "首批素材登记", 2)
    add_table(
        doc,
        ["素材ID", "类型／文件名", "来源／权利人", "授权范围", "审核人／日期"],
        [("", "", "", "", "") for _ in range(3)],
        [0.9, 1.7, 1.7, 1.45, 1.25],
        font_size=8.2,
        center_cols=(0,),
        zebra=False,
        min_row_mm=9,
    )


def add_week1_deliverables(doc: Document) -> None:
    add_page_title(doc, "07｜第1周四岗位交付清单", "提交截止建议设在周验收会前24小时；证据路径必须指向可打开的文件、原型、记录或数据。")
    add_table(
        doc,
        ["岗位", "本周必须完成", "必须提交的交付物", "验收标准", "验收人"],
        [
            ("1 产品与前端", "参加专业访谈；梳理目标用户、核心流程、P0页面和待决策项。", "P0页面清单V0.1；待决策清单；产品决策记录。", "覆盖首页、体验、结果页和异常入口；待决策项均有责任人、期限和证据。", "指导教师＋岗位2"),
            ("2 3D与技术", "搭建小程序、CloudBase与XR-Frame基础环境；准备6台真机清单；比较两条候选路线技术差异。", "可运行工程骨架；XR基础场景；设备清单；技术可行性记录。", "至少1台真机运行XR基础场景；明确模型、手势、性能和降级风险。", "指导教师＋岗位1"),
            ("3 内容与档案", "组织首次访谈；核实对象信息；明确可拍摄内容、公开边界和授权范围。", "访谈纪要；专业确认；授权事项清单；首批素材目录。", "工艺主线有专业依据，或清楚标为待确认；每类素材有授权状态。", "指导教师"),
            ("4 调研与实践", "定义一般用户层与正式研究层目标、样本口径、招募路径和前后测构念。", "研究目标说明；样本口径表；招募计划；前后测构念清单。", "区分800用户、100配对样本与30—50人测试，并写清排除规则和证据来源。", "指导教师"),
        ],
        [1.05, 1.55, 1.55, 2.15, 0.7],
        font_size=7.9,
        center_cols=(0, 4),
        min_row_mm=18,
    )
    add_heading(doc, "提交与命名", 2)
    add_table(
        doc,
        ["项目", "统一要求"],
        [
            ("文件命名", "岗位号_交付物名称_版本_日期，例如：03_首次访谈纪要_V0.1_20260722.docx"),
            ("证据位置", "共享文件夹或仓库中的可访问路径；不得仅在个人聊天记录或本地桌面保存。"),
            ("版本交接", "保留源文件、导出文件、修改人、修改日期和已知问题。"),
            ("状态口径", "未开始、进行中、待验收、已通过、需返工、已阻塞。"),
        ],
        [1.3, 5.7],
        font_size=8.8,
        center_cols=(0,),
        zebra=False,
    )


def add_acceptance_record(doc: Document) -> None:
    add_page_title(doc, "08｜第1周验收记录", "主责人展示交付物，协作者确认接口与可接手性，指导教师记录通过、返工或阻塞结论。")
    add_form_table(
        doc,
        [
            ("验收日期", ""),
            ("验收地点／方式", ""),
            ("参加人员", ""),
            ("记录人", ""),
        ],
        min_row_mm=8,
    )
    add_table(
        doc,
        ["岗位", "交付物与证据位置", "协作者复核", "教师结论", "返工／阻塞期限"],
        [
            ("1 产品与前端", "", "姓名／意见：", "□通过 □返工 □阻塞", ""),
            ("2 3D与技术", "", "姓名／意见：", "□通过 □返工 □阻塞", ""),
            ("3 内容与档案", "", "姓名／意见：", "□通过 □返工 □阻塞", ""),
            ("4 调研与实践", "", "姓名／意见：", "□通过 □返工 □阻塞", ""),
        ],
        [1.15, 2.4, 1.35, 1.3, 0.8],
        font_size=8.2,
        center_cols=(0, 3, 4),
        zebra=False,
        min_row_mm=17,
    )
    add_heading(doc, "问题与闭环", 2)
    add_table(
        doc,
        ["编号", "问题／缺口", "影响", "责任人", "完成期限", "关闭证据／状态"],
        [(str(i), "", "", "", "", "") for i in range(1, 6)],
        [0.5, 2.25, 1.05, 0.85, 0.95, 1.4],
        font_size=8.1,
        center_cols=(0, 2, 3, 4, 5),
        zebra=False,
        min_row_mm=10,
    )
    add_form_table(
        doc,
        [
            ("指导教师结论", "□第1周通过　□有条件通过　□未通过；说明："),
            ("签名／日期", ""),
        ],
        min_row_mm=10,
    )


def add_decision_log(doc: Document) -> None:
    add_page_title(doc, "09｜决策与第2周行动日志", "把访谈结论和周验收决定转化为可追溯行动，避免学生依据口头信息各自理解。")
    add_heading(doc, "关键决策记录", 2)
    add_table(
        doc,
        ["决策项", "当前结论", "证据／确认人", "状态／复核日期"],
        [
            ("首版工艺主线", "□制箔　□贴箔　□仍待确认", "", ""),
            ("三步互动框架", "步骤1：　　　步骤2：　　　步骤3：", "", ""),
            ("临时技术路线", "□已被专业确认替换　□继续内部占位　□停止使用", "", ""),
            ("公开内容边界", "", "", ""),
            ("内容与动作审核人", "", "", ""),
            ("技术降级基线", "□点击／拖动　□2D　□图文　□视频；其他：", "", ""),
            ("研究与传播边界", "研究参与、联系方式与传播激励分开管理。", "", ""),
        ],
        [1.35, 2.65, 1.55, 1.45],
        font_size=8.4,
        center_cols=(0,),
        zebra=False,
        min_row_mm=10,
    )
    add_heading(doc, "第2周行动日志", 2)
    add_table(
        doc,
        ["编号", "行动", "主责", "协作", "期限", "交付证据", "状态"],
        [(str(i), "", "", "", "", "", "未开始") for i in range(1, 7)],
        [0.45, 2.2, 0.75, 0.75, 0.8, 1.4, 0.65],
        font_size=7.9,
        center_cols=(0, 2, 3, 4, 6),
        zebra=False,
        min_row_mm=9,
    )
    add_page_title(
        doc,
        "10｜进入第2周前确认",
        "只有共同闸门和四岗位前置条件均明确，才进入低保真原型、单步骤3D、拍摄方案和研究工具设计。",
    )
    add_table(
        doc,
        ["检查", "进入条件"],
        [
            ("□", "岗位1可以开始低保真主流程，且已标出所有仍待专业确认的内容。"),
            ("□", "岗位2获得至少一个代表性动作与工具材料参考，可启动单步骤3D原型。"),
            ("□", "岗位3拥有拍摄清单、授权边界和可执行的后续沟通安排。"),
            ("□", "岗位4获得页面清单与样本口径，可开始埋点字典和问卷初稿。"),
            ("□", "所有公开内容仍遵守“未审核不发布、未授权不公开”。"),
        ],
        [0.55, 6.45],
        font_size=8.8,
        center_cols=(0,),
        zebra=False,
    )
    add_heading(doc, "启动包归档确认", 2)
    add_form_table(
        doc,
        [
            ("岗位分工表", "文件位置／版本："),
            ("访谈纪要与确认", "文件位置／确认方式："),
            ("授权与素材清单", "文件位置／仍待补证事项："),
            ("第1周验收记录", "结论／返工期限："),
            ("第2周启动时间", "日期／会议安排："),
        ],
        min_row_mm=11,
        font_size=8.8,
    )
    add_callout(
        doc,
        "教师提醒｜",
        "如果专业访谈尚未完成，允许岗位1、2继续低保真和技术可行性工作，但所有文化内容必须标注为待审核，岗位3优先补齐访谈与授权，岗位4不得把目标值写成已取得成果。",
        fill=CREAM,
    )


def validate(doc: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    required = [
        "项目启动会议程",
        "岗位志愿与团队分工",
        "首次专业访谈准备",
        "首次专业访谈提纲",
        "拍摄与授权确认清单",
        "第1周四岗位交付清单",
        "第1周验收记录",
        "决策与第2周行动日志",
        "内部临时技术占位路线",
        "制箔",
        "贴箔",
        "1 产品与前端",
        "2 3D与技术",
        "3 内容与档案",
        "4 调研与实践",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise AssertionError(f"Missing required kickoff content: {missing}")
    forbidden = ["TODO", "TBD", "{{", "}}"]
    hits = [item for item in forbidden if item in text]
    if hits:
        raise AssertionError(f"Template markers found: {hits}")
    if len(doc.tables) < 20:
        raise AssertionError(f"Expected at least 20 operational tables, found {len(doc.tables)}")


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_agenda(doc)
    add_role_selection(doc)
    add_interview_prep(doc)
    add_interview_questions(
        doc,
        "04｜首次专业访谈提纲（上）",
        "先确认对象、首版工艺主线和三步互动，不引导对方接受预设路线。",
        INTERVIEW_PART_1,
    )
    add_interview_questions(
        doc,
        "05｜首次专业访谈提纲（下）",
        "继续确认真实性红线、拍摄与使用范围、审核机制和合作期待。",
        INTERVIEW_PART_2,
    )
    add_interview_conclusion(doc)
    add_authorization(doc)
    add_week1_deliverables(doc)
    add_acceptance_record(doc)
    add_decision_log(doc)
    validate(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
