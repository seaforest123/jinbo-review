from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "gold-foil-heritage" / "首次专业访谈材料包_V0.1_20260719"

FILES = [
    "00_首次专业访谈材料包总目录与使用顺序_V0.1_20260719.docx",
    "01_金箔非遗数字体验项目一页纸_V0.1_20260719.docx",
    "02_三分钟轻量工艺体验纸面原型_V0.1_20260719.docx",
    "03_纸面原型演示脚本与专家反馈表_V0.1_20260719.docx",
    "04_首次专业访谈提纲与60分钟主持脚本_V0.1_20260719.docx",
    "05_拍摄录音及素材使用授权确认单_V0.1_20260719.docx",
    "06_访谈对象身份与联系人信息确认表_V0.1_20260719.docx",
    "07_时间地点确认及微信约访文本_V0.1_20260719.docx",
    "08_现场设备与记录工具检查清单_V0.1_20260719.docx",
    "09_访谈纪要关键结论与会后确认单_V0.1_20260719.docx",
]

REQUIRED = {
    "00": ["9份独立使用文件＋1份总目录", "负责人", "使用时点", "填写要求", "归档位置", "完成状态",
           "访谈前", "访谈中", "访谈后24小时内", "必带材料", "可选材料", "资料完整性", "候选方案／待专业确认"],
    "01": ["广东高校学生", "重点为18—24岁青年", "3分钟轻量工艺体验＋可选深度学习", "互动模拟是核心入口",
           "3D／AR", "五个问题", "身份", "首版路线", "三个互动节点", "真实性红线", "不教授完整真实工艺", "不替代线下传承", "四岗位"],
    "02": ["数字化轻量体验", "制箔候选路线", "贴箔候选路线", "第三种路线", "三步互动", "体验结果",
           "不代表已经掌握真实工艺", "可选深入了解", "人物故事", "工艺知识", "线下活动", "资料页", "待专业确认"],
    "03": ["5—10分钟", "目前尚无可运行Demo", "本次展示的是纸面原型", "两条路线均为候选方案",
           "具体工艺内容待专业确认", "纸面原型版本", "专家第一反应", "不得预设", "专业理由", "真实性风险", "改进优先级", "会后再次确认"],
    "04": ["0—5分钟", "5—15分钟", "15—40分钟", "40—55分钟", "55—60分钟", "第一层必答",
           "第二层", "第三层", "身份称谓", "首版路线", "三个节点", "真实性红线", "记录与使用", "时间不足降级", "纸面原型版本"],
    "05": ["现场录音", "现场录像", "现场拍照", "人物姓名", "公开身份", "肖像", "声音", "工艺过程",
           "工具与材料", "场所", "作品", "团队内部研究", "公开小程序", "攀登计划申报材料", "答辩", "校园讲座", "短视频",
           "授权期限", "署名方式", "审核联系人", "变更／撤回方式", "历史资料", "第三方材料", "不构成授权", "签名"],
    "06": ["常用称呼", "机构", "职务／专业身份", "与金箔工艺的关系", "建议公开表述", "有证据才填写",
           "手机", "微信", "邮箱", "首选渠道", "内容审核联系人", "授权联系人", "紧急联系人", "可公开", "仅内部使用", "不得记录", "确认方式／日期", "后续变更记录"],
    "07": ["首次正式约访", "时间地点确认", "访谈前一天提醒", "临时变更", "重新确认", "60分钟", "预计", "参与人数",
           "纸面原型", "录音录像", "材料线索", "现场再次确认", "05授权确认单", "不等同正式授权"],
    "08": ["负责人", "完成状态", "异常说明", "补救措施", "电量", "存储空间", "充电宝", "录音测试", "备用",
           "静音", "纸面原型", "签字笔", "便签", "夹板", "离线备份", "文件名", "光线", "噪声", "网络", "座位", "拍摄区域", "同意", "资料", "离场前"],
    "09": ["访谈后24小时内", "关联05授权单", "核心观点", "原话摘录", "适用范围／上下文", "路线结论", "专业理由",
           "暂不确定", "三个互动节点", "真实性红线", "公开边界", "审核", "授权摘要", "待补", "已确认", "待确认", "被否定", "需补证据",
           "四岗位第2周行动", "负责人", "截止时间", "交付证据", "简短确认稿", "确认人／方式", "确认日期"],
}

BANNED = [
    "TODO", "TBD", "首版确定采用制箔", "首版确定采用贴箔", "最终选择制箔", "最终选择贴箔",
    "你已经掌握真实工艺", "完整掌握金箔工艺", "无需专业审核", "默认取得授权", "一次签字全部授权",
    "一经签署即同意全部用途",
]


def extract_text(doc: Document) -> str:
    chunks = []
    for section in doc.sections:
        chunks.extend(p.text for p in section.header.paragraphs)
        chunks.extend(p.text for p in section.footer.paragraphs)
    chunks.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def fail(message: str, failures: list[str]) -> None:
    failures.append(message)


def main() -> int:
    failures: list[str] = []
    actual = sorted(p.name for p in OUT_DIR.glob("*.docx")) if OUT_DIR.exists() else []
    if actual != sorted(FILES):
        fail(f"file set mismatch: expected={len(FILES)} actual={len(actual)}", failures)

    texts: dict[str, str] = {}
    for filename in FILES:
        path = OUT_DIR / filename
        code = filename[:2]
        if not path.exists():
            fail(f"missing {filename}", failures)
            continue
        if path.stat().st_size < 35000:
            fail(f"suspiciously small {filename}: {path.stat().st_size}", failures)
        try:
            with zipfile.ZipFile(path) as archive:
                bad = archive.testzip()
                if bad:
                    fail(f"bad ZIP member {filename}: {bad}", failures)
            doc = Document(path)
        except Exception as exc:
            fail(f"cannot open {filename}: {exc}", failures)
            continue
        text = extract_text(doc)
        texts[code] = text
        for phrase in REQUIRED[code]:
            if phrase not in text:
                fail(f"{code} missing phrase: {phrase}", failures)
        for phrase in BANNED:
            if phrase in text:
                fail(f"{code} contains banned phrase: {phrase}", failures)
        if len(doc.tables) < 2:
            fail(f"{code} has too few tables: {len(doc.tables)}", failures)

    if "02" in texts and texts["02"].count("□ 保留　□ 修改　□ 删除　□ 待补证据") < 5:
        fail("02 does not provide the four-way review choice on all five screens", failures)
    if "05" in texts:
        for option, needle in (("同意", "同意"), ("不同意", "不同意"), ("需发布前复核", "发布前复核")):
            if texts["05"].count(needle) < 3:
                fail(f"05 lacks itemized occurrences of {option}", failures)
        if "内部研究" not in texts["05"] or "公开传播" not in texts["05"]:
            fail("05 does not separate internal research from public dissemination", failures)
    if "07" in texts and "只代表参与意愿" not in texts["07"]:
        fail("07 fails to distinguish appointment consent from authorization", failures)
    if "09" in texts and "不扩大05授权确认单的范围" not in texts["09"]:
        fail("09 does not limit its authorization summary to form 05", failures)

    if failures:
        print("FAIL")
        for item in failures:
            print(f"- {item}")
        return 1

    print("PASS")
    print(f"files={len(FILES)}")
    for filename in FILES:
        path = OUT_DIR / filename
        print(f"{filename}\t{path.stat().st_size}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
